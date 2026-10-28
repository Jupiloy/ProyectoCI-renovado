"""Extracción estructural acotada para DOCX, PPTX y XLSX de fuentes del corpus.

No calcula fórmulas, no verifica afirmaciones y no convierte imágenes en texto.
El archivo original y su SHA-256 permanecen como artefacto canónico.
"""

import hashlib
import importlib.metadata
import zipfile
from pathlib import Path

from thesis.models.sources import SourceSegment
from thesis.services.errors import DomainError


MAX_ARCHIVE_MEMBERS = 10_000
MAX_EXPANDED_BYTES = 500 * 1024 * 1024
MAX_MEMBER_BYTES = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200


def _version(package: str) -> str:
    try:
        return importlib.metadata.version(package)
    except importlib.metadata.PackageNotFoundError:
        return "unknown"


def _validate_package(path: Path, expected_main: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_MEMBERS:
                raise DomainError("DOCUMENT_LIMIT", "El paquete contiene demasiados elementos")
            if sum(item.file_size for item in infos) > MAX_EXPANDED_BYTES:
                raise DomainError("DOCUMENT_LIMIT", "El tamaño expandido supera el límite")
            for item in infos:
                if item.file_size > MAX_MEMBER_BYTES or (item.compress_size == 0 and item.file_size > 0):
                    raise DomainError("DOCUMENT_LIMIT", "Un elemento interno supera el límite permitido")
                if item.compress_size and item.file_size / item.compress_size > MAX_COMPRESSION_RATIO:
                    raise DomainError("DOCUMENT_LIMIT", "El paquete excede la razón de compresión permitida")
            required = {"[Content_Types].xml", "_rels/.rels", expected_main}
            if not required.issubset(set(archive.namelist())):
                raise DomainError("INVALID_DOCUMENT", "El archivo no tiene una estructura Office Open XML válida")
    except zipfile.BadZipFile as exc:
        raise DomainError("INVALID_DOCUMENT", "El archivo Office no es un paquete ZIP válido") from exc


def _segment(sequence: int, kind: str, locator: dict, text: str = "", *, method: str,
             version: str, metadata: dict | None = None) -> dict:
    return {
        "sequence_no": sequence, "kind": kind, "locator_json": locator, "text": text,
        "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "extraction_method": method, "extraction_version": version,
        "metadata_json": metadata or {},
    }


def _docx(path: Path) -> tuple[list[dict], dict]:
    _validate_package(path, "word/document.xml")
    from docx import Document

    version = _version("python-docx")
    doc = Document(path)
    segments: list[dict] = []
    paragraph_no = 0
    table_no = 0
    # python-docx exposes body children in original order through iter_inner_content.
    for item in doc.iter_inner_content():
        if hasattr(item, "text") and item.__class__.__name__ == "Paragraph":
            paragraph_no += 1
            value = item.text.strip()
            if value:
                segments.append(_segment(len(segments) + 1, SourceSegment.Kind.PARAGRAPH,
                                         {"paragraph": paragraph_no, "style": item.style.name}, value,
                                         method="PYTHON_DOCX", version=version))
            # Detect pictures embedded in this paragraph without claiming OCR.
            picture_nodes = item._p.xpath(".//wp:docPr")
            for image_index, node in enumerate(picture_nodes, start=1):
                segments.append(_segment(len(segments) + 1, SourceSegment.Kind.IMAGE_REFERENCE,
                                         {"paragraph": paragraph_no, "image_in_paragraph": image_index},
                                         method="PYTHON_DOCX", version=version,
                                         metadata={"image_name": node.get("name", ""), "ocr_performed": False}))
        else:
            table_no += 1
            for row_no, row in enumerate(item.rows, start=1):
                values = [cell.text.strip() for cell in row.cells]
                text = " | ".join(values).strip(" |")
                if text:
                    segments.append(_segment(len(segments) + 1, SourceSegment.Kind.TABLE,
                                             {"table": table_no, "row": row_no,
                                              "cell_count": len(values)}, text,
                                             method="PYTHON_DOCX", version=version,
                                             metadata={"cells": values}))
    return segments, {"paragraphs_seen": paragraph_no, "tables_seen": table_no,
                      "images_ocr": False, "parser": "python-docx", "parser_version": version}


def _pptx(path: Path) -> tuple[list[dict], dict]:
    _validate_package(path, "ppt/presentation.xml")
    from pptx import Presentation

    version = _version("python-pptx")
    presentation = Presentation(path)
    segments: list[dict] = []
    for slide_no, slide in enumerate(presentation.slides, start=1):
        for shape_no, shape in enumerate(slide.shapes, start=1):
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE; avoid enum-version coupling.
                segments.append(_segment(len(segments) + 1, SourceSegment.Kind.IMAGE_REFERENCE,
                                         {"slide": slide_no, "shape": shape_no}, method="PYTHON_PPTX",
                                         version=version, metadata={"image_name": shape.name, "ocr_performed": False}))
            if getattr(shape, "has_table", False):
                for row_no, row in enumerate(shape.table.rows, start=1):
                    values = [cell.text.strip() for cell in row.cells]
                    text = " | ".join(values).strip(" |")
                    if text:
                        segments.append(_segment(len(segments) + 1, SourceSegment.Kind.TABLE,
                                                 {"slide": slide_no, "shape": shape_no, "row": row_no,
                                                  "cell_count": len(values)}, text,
                                                 method="PYTHON_PPTX", version=version,
                                                 metadata={"cells": values}))
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    segments.append(_segment(len(segments) + 1, SourceSegment.Kind.SLIDE_TEXT,
                                             {"slide": slide_no, "shape": shape_no,
                                              "shape_name": shape.name}, text,
                                             method="PYTHON_PPTX", version=version))
    return segments, {"slides_seen": len(presentation.slides), "parser": "python-pptx",
                      "parser_version": version, "images_ocr": False}


def _xlsx(path: Path) -> tuple[list[dict], dict]:
    _validate_package(path, "xl/workbook.xml")
    import openpyxl

    version = _version("openpyxl")
    formulas_book = openpyxl.load_workbook(path, read_only=False, data_only=False, keep_links=False)
    values_book = openpyxl.load_workbook(path, read_only=False, data_only=True, keep_links=False)
    segments: list[dict] = []
    formulas_seen = 0
    images_seen = 0
    try:
        for sheet in formulas_book.worksheets:
            values_sheet = values_book[sheet.title]
            for row in sheet.iter_rows():
                values = []
                cells = []
                for cell in row:
                    formula_or_value = cell.value
                    if formula_or_value is None:
                        continue
                    cached_value = values_sheet[cell.coordinate].value
                    is_formula = isinstance(formula_or_value, str) and formula_or_value.startswith("=")
                    formulas_seen += int(is_formula)
                    values.append(f"{cell.coordinate}: {formula_or_value}")
                    cells.append({"cell": cell.coordinate, "value_or_formula": formula_or_value,
                                  "is_formula": is_formula,
                                  "cached_value": cached_value if is_formula else None,
                                  "cached_value_available": cached_value is not None if is_formula else None})
                if values:
                    segments.append(_segment(len(segments) + 1, SourceSegment.Kind.SPREADSHEET_ROW,
                                             {"sheet": sheet.title, "row": row[0].row,
                                              "cell_range": f"{row[0].coordinate}:{row[-1].coordinate}"},
                                             " | ".join(values), method="OPENPYXL", version=version,
                                             metadata={"cells": cells, "formula_values_recalculated": False}))
            for image_no, image in enumerate(getattr(sheet, "_images", []), start=1):
                images_seen += 1
                anchor = getattr(image, "anchor", None)
                marker = getattr(anchor, "_from", None)
                locator = {"sheet": sheet.title, "image": image_no}
                if marker is not None:
                    locator.update({"cell_anchor": f"{openpyxl.utils.get_column_letter(marker.col + 1)}{marker.row + 1}"})
                segments.append(_segment(len(segments) + 1, SourceSegment.Kind.IMAGE_REFERENCE,
                                         locator, method="OPENPYXL", version=version,
                                         metadata={"image_format": getattr(image, "format", None), "ocr_performed": False}))
    finally:
        formulas_book.close()
        values_book.close()
    return segments, {"sheets_seen": len(formulas_book.worksheets), "formulas_seen": formulas_seen,
                      "images_seen": images_seen, "formula_values_recalculated": False,
                      "parser": "openpyxl", "parser_version": version, "images_ocr": False}


def extract_office(path: Path, extension: str) -> tuple[list[dict], dict]:
    parsers = {".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx}
    parser = parsers.get(extension.lower())
    if parser is None:
        raise DomainError("UNSUPPORTED_FORMAT", "Formato no admitido")
    segments, details = parser(path)
    if not any(segment["text"].strip() for segment in segments):
        details["extraction_state"] = "NO_TEXT"
    else:
        details["extraction_state"] = "EXTRACTED"
    details["segments_count"] = len(segments)
    return segments, details
