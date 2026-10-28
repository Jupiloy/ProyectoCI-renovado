import hashlib
import base64
from io import BytesIO
from unittest.mock import patch

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import Client, override_settings

from thesis.models import SourceDocument, SourceExcerpt, SourcePassage, SourceSegment
from thesis.services.document_jobs import run_next_document_job
from thesis.services.errors import DomainError
from thesis.services.projects import create_project
from thesis.services.search import index_source_passages, search_passages
from thesis.services.sources import list_source_segments, upload_source_pending


TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p1sAAAAASUVORK5CYII="
)


@pytest.fixture
def context(db, tmp_path):
    user = get_user_model().objects.create_user(
        email="multiformat@example.org", password="test-pass", display_name="Ana",
    )
    outsider = get_user_model().objects.create_user(
        email="multiformat-outsider@example.org", password="test-pass", display_name="Otra",
    )
    return user, outsider, create_project(user, "Corpus multiformato"), tmp_path


def _docx() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph("El estudio describe el método de validación.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Variable", "Indicador"
    table.cell(1, 0).text, table.cell(1, 1).text = "Usabilidad", "Tareas completadas"
    doc.add_picture(BytesIO(TINY_PNG))
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.enum.text import PP_ALIGN
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Resultados de la evaluación"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Métrica", "Valor"
    table.cell(1, 0).text, table.cell(1, 1).text = "Tiempo", "12 min"
    slide.shapes.add_picture(BytesIO(TINY_PNG), Inches(6), Inches(1))
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Resultados piloto"
    sheet["A1"], sheet["B1"] = "Casos", "Éxitos"
    sheet["A2"], sheet["B2"] = 10, 8
    sheet["C2"] = "=B2/A2"
    from openpyxl.drawing.image import Image
    sheet.add_image(Image(BytesIO(TINY_PNG)), "D4")
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(("filename", "builder", "expected_kind"), [
    ("evidencia.docx", _docx, "TABLE"),
    ("resultados.pptx", _pptx, "TABLE"),
    ("datos.xlsx", _xlsx, "SPREADSHEET_ROW"),
])
def test_office_upload_extracts_native_locators_and_indexes(context, filename, builder, expected_kind):
    user, _, project, media_root = context
    raw = builder()
    with override_settings(MEDIA_ROOT=media_root), patch("thesis.services.search.embed_texts") as embed:
        embed.side_effect = lambda values: [[0.01] * 384 for _ in values]
        pending = upload_source_pending(user, project.id, filename, SimpleUploadedFile(filename, raw))
        assert pending["upload_status"] == "EXTRACTION_QUEUED"
        assert pending["verification_status"] == "PENDING_EXTRACTION"
        assert pending["full_text_extracted"] is False
        assert run_next_document_job() is True

        source = SourceDocument.objects.get(pk=pending["source_id"])
        assert source.original_filename == filename
        original = media_root / source.storage_key
        assert original.read_bytes() == raw
        assert hashlib.sha256(original.read_bytes()).hexdigest() == source.sha256
        assert source.verification_status == "TEXT_CHECKED"
        assert source.extraction_state == "EXTRACTED"
        assert source.extraction_details_json["original_preserved"] is True
        assert source.extraction_details_json["claims_verified"] is False
        assert source.segments.filter(kind=expected_kind).exists()
        assert source.segments.filter(kind="IMAGE_REFERENCE").exists()

        listed = list_source_segments(user, project.id, source.id, limit=100)
        assert listed["total"] == source.segments.count()
        assert all(item["locator"] for item in listed["segments"])

        indexed = index_source_passages(user, project.id, source.id)
        assert indexed["indexed_passages"] >= 1
        results = search_passages(user, project.id, "resultados", limit=5)
        assert results["search_status"] == "CANDIDATES_ONLY"
        assert results["results"]
        assert results["results"][0]["physical_page"] is None
        assert results["results"][0]["segment_id"] is not None
        assert "locator" in results["results"][0]

    assert SourcePassage.objects.filter(segment__source_id=pending["source_id"]).exists()
    assert SourceSegment.objects.filter(source_id=pending["source_id"], text_sha256__isnull=False).exists()


def test_xlsx_records_formula_and_cached_value_without_recalculation(context):
    user, _, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        pending = upload_source_pending(user, project.id, "Planilla", SimpleUploadedFile("datos.xlsx", _xlsx()))
        assert run_next_document_job() is True
    source = SourceDocument.objects.get(pk=pending["source_id"])
    formula_cell = next(cell for segment in source.segments.all() for cell in segment.metadata_json["cells"]
                        if cell["cell"] == "C2")
    assert formula_cell["value_or_formula"] == "=B2/A2"
    assert formula_cell["is_formula"] is True
    assert formula_cell["cached_value_available"] is False
    assert source.extraction_details_json["formula_values_recalculated"] is False


def test_multiformat_upload_rejects_unsupported_and_malformed_package(context, tmp_path):
    user, _, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        with pytest.raises(DomainError) as unsupported:
            upload_source_pending(user, project.id, "CSV", SimpleUploadedFile("datos.csv", b"a,b"))
        assert unsupported.value.code == "UNSUPPORTED_FORMAT"

        pending = upload_source_pending(user, project.id, "No es OOXML", SimpleUploadedFile("mal.docx", b"PKbroken"))
        assert run_next_document_job() is True
        source = SourceDocument.objects.get(pk=pending["source_id"])
        source_job = source.document_jobs.get()
        assert source_job.status == "FAILED"
        assert source_job.error_code == "INVALID_DOCUMENT"


def test_multiformat_source_and_segments_remain_project_scoped(context, tmp_path):
    user, outsider, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        pending = upload_source_pending(user, project.id, "Archivo", SimpleUploadedFile("archivo.docx", _docx()))
        assert run_next_document_job() is True
    with pytest.raises(DomainError) as denied:
        list_source_segments(outsider, project.id, pending["source_id"])
    assert denied.value.code == "NOT_FOUND"

    client = Client()
    client.force_login(user)
    response = client.get(f"/api/v1/projects/{project.id}/sources/{pending['source_id']}/segments")
    assert response.status_code == 200
    assert response.json()["total"] > 0


def test_passage_database_constraint_requires_exactly_one_source_locator(context):
    user, _, project, _ = context
    source = SourceDocument.objects.create(
        project=project, kind=SourceDocument.Kind.UPLOAD, record_provider="",
        metadata_origin=SourceDocument.MetadataOrigin.USER_UPLOADED_FILE,
        title="Localizador", storage_key="sources/example.docx", sha256="a" * 64,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        imported_by=user,
    )
    excerpt = SourceExcerpt.objects.create(
        source=source, chunk_no=1, page_no=1, text="Página original",
        text_sha256=hashlib.sha256("Página original".encode("utf-8")).hexdigest(),
    )
    segment = SourceSegment.objects.create(
        source=source, sequence_no=1, kind=SourceSegment.Kind.PARAGRAPH,
        locator_json={"paragraph": 1}, text="Párrafo original",
        text_sha256=hashlib.sha256("Párrafo original".encode("utf-8")).hexdigest(),
        extraction_method="PYTHON_DOCX", extraction_version="test",
    )
    common = {"start_char": 0, "text": "pasaje", "text_sha256": hashlib.sha256(b"pasaje").hexdigest(),
              "embedding": [0.0] * 384, "embedding_model": "test"}
    with pytest.raises(IntegrityError), transaction.atomic():
        SourcePassage.objects.create(**common)
    with pytest.raises(IntegrityError), transaction.atomic():
        SourcePassage.objects.create(**common, excerpt=excerpt, segment=segment)
