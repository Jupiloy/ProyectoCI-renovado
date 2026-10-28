import hashlib
import json
from io import BytesIO
from uuid import uuid4

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError, transaction
from django.test import Client, override_settings

from thesis.models import (ResearchDataset, ResearchDatasetVersion, ResearchTableRecipeVersion,
                           ResearchFigureVersion)
from thesis.services.errors import DomainError
from thesis.services.projects import create_project
from thesis.services.research_data import (
    add_dataset_version, create_dataset, get_dataset_file, list_dataset_versions,
)
from thesis.services.research_tables import (
    create_table_recipe, list_table_recipe_versions,
)
from thesis.services.research_figures import create_figure, get_figure_preview


@pytest.fixture
def context(db, tmp_path):
    owner = get_user_model().objects.create_user(
        email="dataset-owner@example.org", password="test-pass", display_name="Titular",
    )
    outsider = get_user_model().objects.create_user(
        email="dataset-outsider@example.org", password="test-pass", display_name="Externo",
    )
    return owner, outsider, create_project(owner, "Datos de prueba"), tmp_path


def csv_file(name="resultados.csv", body="grupo,valor\nA,2\nB,4\n"):
    return SimpleUploadedFile(name, body.encode("utf-8"), content_type="text/csv")


def xlsx_file(name="resultados.xlsx"):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mediciones"
    sheet.append(["grupo", "valor"])
    sheet.append(["A", 2])
    sheet.append(["B", 4])
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue(),
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_csv_upload_versions_original_hash_schema_and_download(context):
    owner, _, project, media_root = context
    key = uuid4()
    raw = b"grupo,valor\nA,2\nB,4\n"
    with override_settings(MEDIA_ROOT=media_root):
        first = create_dataset(owner, project.id, "Resultados", csv_file(), key)
        second = add_dataset_version(owner, project.id, first["dataset"]["id"],
                                     csv_file(body="grupo,valor\nA,3\nB,5\n"), uuid4(), 1)
        version_one, path_one = get_dataset_file(owner, project.id, first["dataset"]["id"], 1)

    assert first["created"] is True
    assert first["version"]["sha256"] == hashlib.sha256(raw).hexdigest()
    assert first["version"]["row_count"] == 2
    assert first["version"]["column_count"] == 2
    assert first["version"]["columns"][0]["observed_type"] == "TEXT"
    assert first["version"]["columns"][1]["observed_type"] == "INTEGER"
    assert second["version"]["version_no"] == 2
    assert version_one.sha256 == hashlib.sha256(path_one.read_bytes()).hexdigest()
    assert path_one.read_bytes() == raw


def test_same_idempotency_key_and_content_returns_existing_version(context):
    owner, _, project, media_root = context
    key = uuid4()
    with override_settings(MEDIA_ROOT=media_root):
        first = create_dataset(owner, project.id, "Resultados", csv_file(), key)
        retry = create_dataset(owner, project.id, "Nombre ignorado al reintentar", csv_file(), key)
    assert retry["created"] is False
    assert retry["idempotent"] is True
    assert retry["dataset"]["id"] == first["dataset"]["id"]
    assert ResearchDataset.objects.count() == 1
    assert ResearchDatasetVersion.objects.count() == 1


def test_csv_invalid_headers_or_shape_are_rejected_without_persisting(context):
    owner, _, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root), pytest.raises(DomainError) as duplicate:
        create_dataset(owner, project.id, "Duplicadas", csv_file(body="A,a\n1,2\n"), uuid4())
    assert duplicate.value.code == "INVALID_DATASET"
    with override_settings(MEDIA_ROOT=media_root), pytest.raises(DomainError) as wrong_shape:
        create_dataset(owner, project.id, "Forma", csv_file(body="A,B\n1,2,3\n"), uuid4())
    assert wrong_shape.value.code == "INVALID_DATASET"
    assert ResearchDataset.objects.count() == 0


def test_xlsx_requires_sheet_for_multi_sheet_and_reports_formula_caveat(context):
    owner, _, project, media_root = context
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    sheet.append(["casos", "ratio"])
    sheet.append([10, "=5/10"])
    workbook.create_sheet("Notas")
    buffer = BytesIO()
    workbook.save(buffer)
    raw_file = SimpleUploadedFile("datos.xlsx", buffer.getvalue())

    with override_settings(MEDIA_ROOT=media_root), pytest.raises(DomainError) as missing_sheet:
        create_dataset(owner, project.id, "Libro", raw_file, uuid4())
    assert missing_sheet.value.code == "VALIDATION_ERROR"
    with override_settings(MEDIA_ROOT=media_root):
        result = create_dataset(owner, project.id, "Libro", xlsx_file_with_formula(), uuid4(), sheet_name="Mediciones")
    metadata = result["version"]["import_metadata"]
    assert metadata["selected_sheet"] == "Mediciones"
    assert metadata["formula_values_recalculated"] is False
    assert metadata["formula_cells"] == 1
    assert metadata["formula_cells_without_cached_value"] == 1
    with override_settings(MEDIA_ROOT=media_root):
        table = create_table_recipe(
            owner, project.id, result["dataset"]["id"], 1, "Valores guardados",
            {"columns": ["casos", "ratio"], "limit": 10}, uuid4(),
        )
    assert table["version"]["result"]["rows"] == [{"casos": 10, "ratio": None}]


def xlsx_file_with_formula():
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mediciones"
    sheet.append(["casos", "ratio"])
    sheet.append([10, "=5/10"])
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile("mediciones.xlsx", buffer.getvalue())


def test_dataset_write_permission_scope_and_stale_version(context):
    owner, outsider, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        created = create_dataset(owner, project.id, "Datos", csv_file(), uuid4())
        with pytest.raises(DomainError) as denied:
            create_dataset(outsider, project.id, "No autorizado", csv_file(), uuid4())
        assert denied.value.code == "NOT_FOUND"
        with pytest.raises(DomainError) as stale:
            add_dataset_version(owner, project.id, created["dataset"]["id"], csv_file(), uuid4(), 2)
        assert stale.value.code == "VERSION_CONFLICT"


def test_dataset_api_upload_list_versions_and_download(context):
    owner, _, project, media_root = context
    client = Client()
    client.force_login(owner)
    with override_settings(MEDIA_ROOT=media_root):
        response = client.post(
            f"/api/v1/projects/{project.id}/datasets",
            {"name": "Datos", "description": "Piloto", "idempotency_key": str(uuid4()), "file": csv_file()},
        )
        assert response.status_code == 201, response.content
        dataset_id = response.json()["dataset"]["id"]
        listed = client.get(f"/api/v1/projects/{project.id}/datasets")
        assert listed.status_code == 200
        assert listed.json()["datasets"][0]["id"] == dataset_id
        versions = client.get(f"/api/v1/projects/{project.id}/datasets/{dataset_id}/versions")
        assert versions.json()["versions"][0]["version_no"] == 1
        downloaded = client.get(f"/api/v1/projects/{project.id}/datasets/{dataset_id}/versions/1/download")
        assert downloaded.status_code == 200
        assert b"grupo,valor" in b"".join(downloaded.streaming_content)


@pytest.mark.django_db(transaction=True)
def test_dataset_versions_are_append_only(context):
    owner, _, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        created = create_dataset(owner, project.id, "Inmutable", csv_file(), uuid4())
    version_id = created["version"]["id"]
    with pytest.raises(DatabaseError):
        with transaction.atomic():
            ResearchDatasetVersion.objects.filter(pk=version_id).update(row_count=99)
    with pytest.raises(DatabaseError):
        with transaction.atomic():
            ResearchDatasetVersion.objects.filter(pk=version_id).delete()


def test_version_list_is_scoped_to_project(context):
    owner, outsider, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        created = create_dataset(owner, project.id, "Datos privados", csv_file(), uuid4())
    with pytest.raises(DomainError) as denied:
        list_dataset_versions(outsider, project.id, created["dataset"]["id"])
    assert denied.value.code == "NOT_FOUND"


def test_table_recipe_filters_sorts_projects_and_hashes_reproducible_output(context):
    owner, _, project, media_root = context
    body = "grupo,valor,nota\nA,2,primero\nB,10,segundo\nC,4,tercero\n"
    with override_settings(MEDIA_ROOT=media_root):
        created = create_dataset(owner, project.id, "Resultados", csv_file(body=body), uuid4())
        result = create_table_recipe(
            owner, project.id, created["dataset"]["id"], 1, "Valores válidos",
            {"columns": ["grupo", "valor"], "filters": [{"column": "valor", "operator": "GTE", "value": 4}],
             "sort": [{"column": "valor", "direction": "DESC"}], "limit": 5}, uuid4(),
        )
    table = result["version"]["result"]
    assert result["created"] is True
    assert table["rows"] == [{"grupo": "B", "valor": "10"}, {"grupo": "C", "valor": "4"}]
    assert table["matched_row_count"] == 2
    assert table["truncated"] is False
    assert result["version"]["result_sha256"] == hashlib.sha256(
        json.dumps(table, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def test_table_recipe_versions_are_immutable_and_idempotent(context):
    owner, _, project, media_root = context
    key = uuid4()
    spec = {"columns": ["grupo", "valor"], "limit": 10}
    with override_settings(MEDIA_ROOT=media_root):
        created = create_dataset(owner, project.id, "Resultados", csv_file(), uuid4())
        first = create_table_recipe(owner, project.id, created["dataset"]["id"], 1,
                                    "Resumen", spec, key)
        retry = create_table_recipe(owner, project.id, created["dataset"]["id"], 1,
                                    "Resumen", spec, key, expected_version=1)
        second = create_table_recipe(owner, project.id, created["dataset"]["id"], 1,
                                     "Resumen", {"columns": ["valor"], "limit": 1}, uuid4(),
                                     expected_version=1)
        history = list_table_recipe_versions(owner, project.id, created["dataset"]["id"],
                                             first["recipe"]["id"])
    assert retry["idempotent"] is True
    assert retry["version"]["id"] == first["version"]["id"]
    assert second["version"]["version_no"] == 2
    assert [item["version_no"] for item in history["versions"]] == [2, 1]
    assert history["versions"][1]["result"] == first["version"]["result"]


def test_table_recipe_rejects_free_code_unknown_fields_and_stale_expected_version(context):
    owner, _, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        created = create_dataset(owner, project.id, "Resultados", csv_file(), uuid4())
        base = {"columns": ["valor"], "limit": 10}
        with pytest.raises(DomainError) as unknown:
            create_table_recipe(owner, project.id, created["dataset"]["id"], 1, "Mala",
                                {**base, "python": "__import__('os').system('id')"}, uuid4())
        assert unknown.value.code == "VALIDATION_ERROR"
        with pytest.raises(DomainError) as invalid_operator:
            create_table_recipe(owner, project.id, created["dataset"]["id"], 1, "Comparación textual",
                                {"columns": ["grupo"], "filters": [
                                    {"column": "grupo", "operator": "GT", "value": "A"}], "limit": 10}, uuid4())
        assert invalid_operator.value.code == "VALIDATION_ERROR"
        first = create_table_recipe(owner, project.id, created["dataset"]["id"], 1, "Control", base, uuid4())
        with pytest.raises(DomainError) as stale:
            create_table_recipe(owner, project.id, created["dataset"]["id"], 1, "Control",
                                base, uuid4(), expected_version=2)
        assert stale.value.code == "VERSION_CONFLICT"
        with pytest.raises(DomainError) as denied:
            create_table_recipe(get_user_model().objects.create_user(
                email="table-outsider@example.org", password="pw", display_name="Externo",
            ), project.id, created["dataset"]["id"], 1, "No", base, uuid4())
        assert denied.value.code == "NOT_FOUND"
        assert first["version"]["version_no"] == 1


def test_figure_rendering_versions_hashes_idempotency_and_truncation_confirmation(context):
    owner, outsider, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        dataset = create_dataset(owner, project.id, "Datos de figuras", csv_file(), uuid4())
        table = create_table_recipe(
            owner, project.id, dataset["dataset"]["id"], 1, "Tabla truncada",
            {"columns": ["grupo", "valor"], "limit": 1}, uuid4(),
        )
        figure_input = {"chart_type": "BAR", "x_column": "grupo", "y_column": "valor",
                        "title": "Valores observados", "x_label": "Grupo", "y_label": "Valor",
                        "alt_text": "Gráfico de barras: A vale 2."}
        with pytest.raises(DomainError) as consent:
            create_figure(owner, project.id, dataset["dataset"]["id"], table["recipe"]["id"], 1,
                          "Figura 1", figure_input, uuid4())
        assert consent.value.code == "TRUNCATED_INPUT_CONFIRMATION_REQUIRED"
        key = uuid4()
        first = create_figure(owner, project.id, dataset["dataset"]["id"], table["recipe"]["id"], 1,
                              "Figura 1", figure_input, key, accept_truncated_input=True)
        retry = create_figure(owner, project.id, dataset["dataset"]["id"], table["recipe"]["id"], 1,
                              "Figura 1", figure_input, key, expected_version=1, accept_truncated_input=True)
        preview = get_figure_preview(owner, project.id, dataset["dataset"]["id"],
                                     first["figure"]["id"], 1)
        data = preview.read_bytes()
        assert data.startswith(b"\x89PNG\r\n\x1a\n")
        assert hashlib.sha256(data).hexdigest() == first["version"]["output_sha256"]
        assert first["version"]["accessibility"]["source_truncated"] is True
        assert retry["idempotent"] is True
        assert retry["version"]["id"] == first["version"]["id"]
        second = create_figure(owner, project.id, dataset["dataset"]["id"], table["recipe"]["id"], 1,
                               "Figura 1", {**figure_input, "title": "Valores revisados"}, uuid4(),
                               expected_version=1, accept_truncated_input=True)
        assert second["version"]["version_no"] == 2
        assert ResearchFigureVersion.objects.count() == 2
        with pytest.raises(DomainError) as denied:
            get_figure_preview(outsider, project.id, dataset["dataset"]["id"], first["figure"]["id"], 1)
        assert denied.value.code == "NOT_FOUND"


def test_figure_rejects_non_numeric_y_and_unknown_fields(context):
    owner, _, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        dataset = create_dataset(owner, project.id, "Datos", csv_file(), uuid4())
        table = create_table_recipe(owner, project.id, dataset["dataset"]["id"], 1, "Tabla",
                                    {"columns": ["grupo", "valor"], "limit": 10}, uuid4())
        spec = {"chart_type": "BAR", "x_column": "grupo", "y_column": "grupo",
                "title": "Inválido", "x_label": "X", "y_label": "Y", "alt_text": "Texto alternativo",
                "python": "no ejecutar"}
        with pytest.raises(DomainError) as invalid:
            create_figure(owner, project.id, dataset["dataset"]["id"], table["recipe"]["id"], 1,
                          "Figura", spec, uuid4())
        assert invalid.value.code == "VALIDATION_ERROR"


def test_research_figure_api_create_list_metadata_and_preview(context):
    owner, _, project, media_root = context
    client = Client()
    client.force_login(owner)
    with override_settings(MEDIA_ROOT=media_root):
        dataset = create_dataset(owner, project.id, "API datos", csv_file(), uuid4())
        table = create_table_recipe(owner, project.id, dataset["dataset"]["id"], 1, "API tabla",
                                    {"columns": ["grupo", "valor"], "limit": 10}, uuid4())
        collection = (f"/api/v1/projects/{project.id}/datasets/{dataset['dataset']['id']}"
                      f"/table-recipes/{table['recipe']['id']}/versions/1/figures")
        payload = {"name": "API figura", "spec": {
            "chart_type": "BAR", "x_column": "grupo", "y_column": "valor", "title": "API",
            "x_label": "Grupo", "y_label": "Valor", "alt_text": "Barras de API.",
        }, "idempotency_key": str(uuid4())}
        created = client.post(collection, data=json.dumps(payload), content_type="application/json")
        assert created.status_code == 201, created.content
        figure_id = created.json()["figure"]["id"]
        listed = client.get(collection)
        assert listed.status_code == 200
        assert listed.json()["figures"][0]["id"] == figure_id
        metadata = client.get(
            f"/api/v1/projects/{project.id}/datasets/{dataset['dataset']['id']}"
            f"/figures/{figure_id}/versions/1"
        )
        assert metadata.status_code == 200
        assert metadata.json()["version"]["mime_type"] == "image/png"
        preview = client.get(
            f"/api/v1/projects/{project.id}/datasets/{dataset['dataset']['id']}"
            f"/figures/{figure_id}/versions/1/preview.png"
        )
        assert preview.status_code == 200
        assert preview["Content-Type"] == "image/png"
        assert b"".join(preview.streaming_content).startswith(b"\x89PNG\r\n\x1a\n")


@pytest.mark.django_db(transaction=True)
def test_figure_versions_are_append_only(context):
    owner, _, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        dataset = create_dataset(owner, project.id, "Datos", csv_file(), uuid4())
        table = create_table_recipe(owner, project.id, dataset["dataset"]["id"], 1, "Tabla",
                                    {"columns": ["grupo", "valor"], "limit": 10}, uuid4())
        result = create_figure(owner, project.id, dataset["dataset"]["id"], table["recipe"]["id"], 1,
                               "Figura", {"chart_type": "BAR", "x_column": "grupo", "y_column": "valor",
                                          "title": "Título", "x_label": "Grupo", "y_label": "Valor",
                                          "alt_text": "Barras por grupo"}, uuid4())
    with pytest.raises(DatabaseError):
        with transaction.atomic():
            ResearchFigureVersion.objects.filter(pk=result["version"]["id"]).update(size_bytes=1)
    with pytest.raises(DatabaseError):
        with transaction.atomic():
            ResearchFigureVersion.objects.filter(pk=result["version"]["id"]).delete()


def test_table_recipe_api_and_integrity_check_on_read(context):
    owner, _, project, media_root = context
    client = Client()
    client.force_login(owner)
    with override_settings(MEDIA_ROOT=media_root):
        uploaded = client.post(
            f"/api/v1/projects/{project.id}/datasets",
            {"name": "Datos", "idempotency_key": str(uuid4()), "file": csv_file()},
        )
        dataset_id = uploaded.json()["dataset"]["id"]
        response = client.post(
            f"/api/v1/projects/{project.id}/datasets/{dataset_id}/versions/1/table-recipes",
            data=json.dumps({"name": "Tabla", "spec": {"columns": ["grupo"], "limit": 10},
                             "idempotency_key": str(uuid4())}),
            content_type="application/json",
        )
        assert response.status_code == 201, response.content
        recipe_id = response.json()["recipe"]["id"]
        read = client.get(
            f"/api/v1/projects/{project.id}/datasets/{dataset_id}/table-recipes/{recipe_id}/versions/1"
        )
        assert read.status_code == 200
        assert read.json()["version"]["result"]["row_count"] == 2
    assert read.json()["version"]["result"]["row_count"] == 2


@pytest.mark.django_db(transaction=True)
def test_table_recipe_versions_are_append_only(context):
    owner, _, project, media_root = context
    with override_settings(MEDIA_ROOT=media_root):
        dataset = create_dataset(owner, project.id, "Datos", csv_file(), uuid4())
        result = create_table_recipe(owner, project.id, dataset["dataset"]["id"], 1,
                                     "Inmutable", {"columns": ["grupo"], "limit": 10}, uuid4())
    recipe_version = ResearchTableRecipeVersion.objects.get(pk=result["version"]["id"])
    with pytest.raises(DatabaseError):
        with transaction.atomic():
            ResearchTableRecipeVersion.objects.filter(pk=recipe_version.pk).update(row_count=99)
    with pytest.raises(DatabaseError):
        with transaction.atomic():
            ResearchTableRecipeVersion.objects.filter(pk=recipe_version.pk).delete()
