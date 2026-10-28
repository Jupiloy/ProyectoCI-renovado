"""Pruebas exhaustivas para la instantánea integral e inmutable de capítulo (B1.1)."""

import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from uuid import uuid4

import pytest
from docx import Document
from django.contrib.auth import get_user_model
from django.db import IntegrityError, close_old_connections, connection, transaction
from django.db.utils import DatabaseError, ProgrammingError
from django.test import Client

from thesis.models import (
    Chapter,
    ChapterRevision,
    ChapterDocumentRevision,
    ChapterSnapshot,
    ResearchProfile,
    ReviewComment,
    ReviewEvent,
    SourceDocument,
    SourceExcerpt,
    EvidenceLink,
    StepProgress,
    WorkflowTaskExecution,
    ResearchFigureVersion,
)
from thesis.services.errors import DomainError
from thesis.services.projects import create_project
from thesis.services.reviews import add_comment, assign_advisor, decide_chapter
from thesis.services.steps import (
    get_chapter,
    get_chapter_snapshot,
    list_chapter_snapshots,
    save_step,
    submit_chapter,
)
from thesis.services.exports import export_approved_snapshot_docx
from thesis.services.chapter_documents import save_chapter_document
from thesis.services.research_data import create_dataset
from thesis.services.research_tables import create_table_recipe
from thesis.services.research_figures import create_figure
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from thesis.services.workflows import (
    create_or_update_draft,
    execute_custom_workflow_task,
    get_workflow_state,
    publish_workflow_draft,
)


def _p01_to_p05_outputs():
    candidates = [
        {
            "id": i,
            "titulo_tentativo": f"Tema candidato automatizado número {i}",
            "pros": ["Acceso a datos primarios"],
            "contras": ["Tiempo limitado"],
            "viabilidad_preliminar": "MEDIA",
        }
        for i in range(1, 6)
    ]
    score = {
        "acceso_a_datos": 4,
        "relevancia_profesional": 4,
        "brecha_conocimiento": 4,
        "recursos_disponibles": 4,
    }
    matrix = [
        {
            "tema_id": i,
            "titulo": f"Tema candidato automatizado número {i}",
            "puntajes": score,
            "puntaje_total": 16,
            "justificacion_eleccion": "Se dispone de datos y tiempo suficiente.",
        }
        for i in range(1, 4)
    ]
    introduction = "\n\n".join(["palabra " * 90] * 3).strip()
    return [
        {"temas_propuestos": candidates, "tema_preseleccionado_id": 1},
        {"matriz_comparativa": matrix, "tema_ganador_id": 1, "tema_ganador_titulo": matrix[0]["titulo"]},
        {
            "ods_principal_numero": 4,
            "ods_principal_nombre": "Educación de calidad",
            "meta_ods_especifica": "4.4",
            "argumento_alineacion": "El trabajo aborda el acceso a formación técnica y competencias profesionales en la comunidad universitaria.",
            "linea_utp_vinculada": "Tecnología educativa",
        },
        {
            "titulo_oficial": "Uso de herramientas digitales en estudiantes de Piura durante 2026",
            "conteo_palabras": 10,
            "variables_confirmadas": ["herramientas digitales"],
            "unidad_analisis_confirmada": "estudiantes",
            "lugar_confirmado": "Piura",
            "periodo_confirmado": 2026,
        },
        {"contenido_markdown": introduction, "num_parrafos": 3, "palabras_totales": 270},
    ]


@pytest.fixture
def snapshot_context(db):
    User = get_user_model()
    owner = User.objects.create_user(email="snap_owner@example.org", password="pwd")
    advisor = User.objects.create_user(email="snap_advisor@example.org", password="pwd")
    outsider = User.objects.create_user(email="snap_outsider@example.org", password="pwd")
    project = create_project(owner, "Proyecto de Tesis Snapshot B1.1")
    assign_advisor(owner, project.id, advisor.email)
    return owner, advisor, outsider, project


# ==============================================================================
# 1. Modo Legado Secuencial: Creación, Contenido, Manifest y Hash Canónico
# ==============================================================================


def test_legacy_sequential_chapter_snapshot_creation_and_integrity(snapshot_context):
    owner, advisor, _, project = snapshot_context
    outputs = _p01_to_p05_outputs()
    codes = ["P01", "P02", "P03", "P04", "P05"]

    # Guardar pasos P01 a P05 para completar Capítulo 1
    for idx, code in enumerate(codes):
        text = f"Texto redactado del paso {code} para la tesis."
        res = save_step(
            actor=owner,
            project_id=project.id,
            step_code=code,
            content_md=text,
            output_data=outputs[idx],
            expected_version=idx,
            idempotency_key=str(uuid4()),
        )
        assert res["step_code"] == code

    p01_revision = ChapterRevision.objects.get(chapter__project=project, step_code="P01")
    source = SourceDocument.objects.create(
        project=project,
        kind=SourceDocument.Kind.EXTERNAL,
        openalex_id="https://openalex.org/W9999999999",
        title="Fuente usada para verificar localización congelada",
        sha256="a" * 64,
    )
    excerpt = SourceExcerpt.objects.create(
        source=source,
        chunk_no=1,
        page_no=7,
        text="Pasaje de prueba con localizador verificable.",
        text_sha256="b" * 64,
    )
    EvidenceLink.objects.create(
        project=project,
        revision=p01_revision,
        excerpt=excerpt,
        claim_text="Afirmación asociada a la fuente.",
        start_char=0,
        end_char=15,
        matched_text="Pasaje de prueba",
        excerpt_sha256=excerpt.text_sha256,
        created_by=owner,
    )

    # Enviar capítulo 1
    submit_res = submit_chapter(owner, project.id, 1)
    assert submit_res["status"] == Chapter.Status.READY
    assert "snapshot" in submit_res
    snap_info = submit_res["snapshot"]
    assert snap_info["snapshot_no"] == 1
    assert snap_info["format_version"] == 1

    # Verificar entidad persistida en base de datos
    chapter = Chapter.objects.get(project=project, number=1)
    assert chapter.status == Chapter.Status.READY
    assert chapter.submitted_snapshot_id is not None

    snapshot = ChapterSnapshot.objects.get(pk=snap_info["id"])
    assert snapshot.chapter == chapter
    assert snapshot.snapshot_no == 1
    assert snapshot.format_version == 1
    project.refresh_from_db()
    assert snapshot.project_title == (project.title or project.name)
    assert snapshot.created_by == owner

    # Verificar que el content_md consolidado contiene todos los textos en orden
    for code in codes:
        assert f"Texto redactado del paso {code}" in snapshot.content_md

    # Verificar el manifest estructurado
    manifest = snapshot.sections_manifest
    assert len(manifest) == 5
    for idx, code in enumerate(codes):
        item = manifest[idx]
        assert item["task_id"] == code
        assert item["legacy_step_code"] == code
        assert item["order_index"] == idx
        assert item["status"] == "COMPLETED"
        assert item["rendered_in_text"] is True
        assert item["revision_id"] is not None
        assert item["content_sha256"] != ""
        assert item["output_json"] == outputs[idx]

    # Verificar recálculo determinista del hash canónico SHA-256
    expected_canonical_payload = {
        "chapter_no": 1,
        "content_md": snapshot.content_md,
        "format_version": 1,
        "project_title": snapshot.project_title,
        "sections_manifest": manifest,
        "state_version": 5,  # lock_version tras 5 guardados de paso
        "workflow_id": None,
        "workflow_content_sha256": "",
        "workflow_version_no": None,
        "document_artifacts_manifest": [],
    }
    expected_bytes = json.dumps(expected_canonical_payload, sort_keys=True, ensure_ascii=False).encode("utf-8")
    expected_hash = hashlib.sha256(expected_bytes).hexdigest()
    assert snapshot.payload_sha256 == expected_hash
    p01_manifest = next(item for item in manifest if item["task_id"] == "P01")
    assert p01_manifest["evidence_links"] == [{
        "id": EvidenceLink.objects.get(revision=p01_revision).id,
        "claim_text": "Afirmación asociada a la fuente.",
        "start_char": 0,
        "end_char": 15,
        "matched_text": "Pasaje de prueba",
        "excerpt_id": excerpt.id,
        "excerpt_sha256": excerpt.text_sha256,
        "excerpt__page_no": 7,
        "excerpt__source_id": str(source.id),
        "excerpt__source__sha256": source.sha256,
        "excerpt__source__title": source.title,
        "status": EvidenceLink.Status.CANDIDATE,
    }]

    # Verificar registro en ReviewEvent con puntero al snapshot
    event = ReviewEvent.objects.filter(chapter=chapter, action=ReviewEvent.Action.SUBMIT).first()
    assert event is not None
    assert event.snapshot == snapshot
    assert event.actor == owner


# ==============================================================================
# 2. Modo Configurable: Orden de Workflow, Opcional Omitida y Tareas Custom USR_
# ==============================================================================


def test_configurable_graph_chapter_snapshot_with_custom_and_optional_tasks(snapshot_context):
    owner, advisor, _, project = snapshot_context
    outputs = _p01_to_p05_outputs()

    # Activar workflow configurable
    get_workflow_state(owner, project.id)
    draft_tasks = [
        {"task_id": "P01", "chapter_no": 1, "is_optional": False, "order_index": 10},
        {"task_id": "P02", "chapter_no": 1, "is_optional": False, "order_index": 20},
        {"task_id": "P03", "chapter_no": 1, "is_optional": False, "order_index": 30},
        {"task_id": "P04", "chapter_no": 1, "is_optional": False, "order_index": 40},
        # P05 marcada como opcional
        {"task_id": "P05", "chapter_no": 1, "is_optional": True, "order_index": 50},
        # Tarea custom USR_SURVEY en capítulo 1
        {
            "task_id": "USR_SURVEY",
            "title": "Encuesta Diagnóstica de Campo",
            "chapter_no": 1,
            "is_optional": False,
            "order_index": 45,
            "prompt_text": "Analizar respuestas de encuesta diagnóstica",
        },
    ]
    create_or_update_draft(owner, project.id, {"tasks": draft_tasks}, expected_version=0)
    publish_workflow_draft(owner, project.id, "Publicar recorrido con USR_SURVEY", expected_version=1)

    profile = ResearchProfile.objects.get(project=project)
    active_wf = profile.active_workflow
    prior_wf_version_no = max(0, active_wf.version_no - 1)

    # Completar P01 a P04 (P05 queda sin completar, al ser opcional)
    saved_revisions = {}
    for idx, code in enumerate(["P01", "P02", "P03", "P04"]):
        saved_revisions[code] = save_step(
            owner,
            project.id,
            code,
            f"Contenido configurable de {code}",
            outputs[idx],
            expected_version=profile.lock_version,
            idempotency_key=str(uuid4()),
        )
        profile.refresh_from_db()

    # Una revisión histórica de otra versión del recorrido no debe colarse en el snapshot.
    ChapterRevision.objects.create(
        chapter=Chapter.objects.get(project=project, number=1),
        revision_no=100,
        step_code="P01",
        idempotency_key=uuid4(),
        content_md="Contenido obsoleto de workflow anterior que no debe aparecer.",
        content_sha256="c" * 64,
        workflow=active_wf,
        workflow_version_no=prior_wf_version_no,
        created_by=owner,
    )

    # Ejecutar la tarea custom obligatoria USR_SURVEY
    survey_output = {"n_encuestados": 45, "hallazgo_clave": "78% requiere soporte"}
    exec_res = execute_custom_workflow_task(
        actor=owner,
        project_id=project.id,
        task_id="USR_SURVEY",
        input_data={"descripcion": "Encuesta diagnóstica aplicada a estudiantes"},
        output_data=survey_output,
        idempotency_key=uuid4(),
    )
    assert exec_res["task_id"] == "USR_SURVEY"

    # Enviar capítulo 1
    submit_res = submit_chapter(owner, project.id, 1)
    assert submit_res["status"] == Chapter.Status.READY
    snap_id = submit_res["snapshot"]["id"]

    snapshot = ChapterSnapshot.objects.get(pk=snap_id)
    manifest = snapshot.sections_manifest

    # Debe contener 6 tareas: P01(10), P02(20), P03(30), P04(40), USR_SURVEY(45), P05(50)
    assert len(manifest) == 6
    task_order = [item["task_id"] for item in manifest]
    assert task_order == ["P01", "P02", "P03", "P04", "USR_SURVEY", "P05"]

    # Inspeccionar entrada USR_SURVEY en el manifest
    usr_entry = next(item for item in manifest if item["task_id"] == "USR_SURVEY")
    assert usr_entry["status"] == "COMPLETED"
    assert usr_entry["execution_id"] == exec_res["execution_id"]
    assert usr_entry["output_json"] == survey_output
    assert usr_entry["rendered_in_text"] is False
    assert "approved" not in usr_entry  # No inventa estado "aprobado"
    assert usr_entry["workflow_id"] == active_wf.pk
    assert usr_entry["workflow_version_no"] == active_wf.version_no
    assert usr_entry["input_json"] == {"descripcion": "Encuesta diagnóstica aplicada a estudiantes"}
    assert usr_entry["prompt_version_used"] is not None
    assert usr_entry["created_by_id"] == owner.pk

    p01_entry = next(item for item in manifest if item["task_id"] == "P01")
    assert p01_entry["revision_id"] == saved_revisions["P01"]["revision_id"]
    assert "Contenido obsoleto" not in snapshot.content_md
    submit_event = ReviewEvent.objects.get(chapter=snapshot.chapter, action=ReviewEvent.Action.SUBMIT)
    assert submit_event.revision_id == saved_revisions["P04"]["revision_id"]

    # Inspeccionar entrada P05 (opcional omitida)
    p05_entry = next(item for item in manifest if item["task_id"] == "P05")
    assert p05_entry["status"] == "SKIPPED"
    assert p05_entry["is_optional"] is True
    assert p05_entry["rendered_in_text"] is False


# ==============================================================================
# 3. Idempotencia en READY y Nueva Instantánea tras Observación y Corrección
# ==============================================================================


def test_idempotent_submit_in_ready_and_new_snapshot_after_changes(snapshot_context):
    owner, advisor, _, project = snapshot_context
    outputs = _p01_to_p05_outputs()

    for idx, code in enumerate(["P01", "P02", "P03", "P04", "P05"]):
        save_step(owner, project.id, code, f"Borrador {code}", outputs[idx], idx, str(uuid4()))

    # Primer envío
    res1 = submit_chapter(owner, project.id, 1)
    snap1_id = res1["snapshot"]["id"]
    assert res1["snapshot"]["snapshot_no"] == 1
    assert ChapterSnapshot.objects.filter(chapter__project=project, chapter__number=1).count() == 1

    # Reintento de envío mientras el capítulo sigue en READY -> Es idempotente
    repeat_res = submit_chapter(owner, project.id, 1)
    assert repeat_res["snapshot"]["id"] == snap1_id
    assert repeat_res["snapshot"]["snapshot_no"] == 1
    # No se crea una segunda fila
    assert ChapterSnapshot.objects.filter(chapter__project=project, chapter__number=1).count() == 1

    # Asesor observa el capítulo
    chapter = Chapter.objects.get(project=project, number=1)
    current_snapshot = ChapterSnapshot.objects.get(pk=snap1_id)
    selected_text = current_snapshot.content_md[:8]
    add_comment(
        advisor,
        project.id,
        1,
        start_cp=0,
        end_cp=len(selected_text),
        selected_text=selected_text,
        body="Mejorar redacción de este párrafo",
        snapshot_id=snap1_id,
    )
    decide_chapter(advisor, project.id, 1, ReviewEvent.Action.REQUEST_CHANGES)
    chapter.refresh_from_db()
    assert chapter.status == Chapter.Status.CHANGES_REQUESTED

    # El tesista guarda una corrección en P05
    profile = ResearchProfile.objects.get(project=project)
    save_step(
        owner,
        project.id,
        "P05",
        "Texto corregido y mejorado para P05 conforme a observaciones.",
        outputs[4],
        profile.lock_version,
        str(uuid4()),
    )

    # El tesista vuelve a enviar el capítulo
    res2 = submit_chapter(owner, project.id, 1)
    snap2_id = res2["snapshot"]["id"]
    assert res2["snapshot"]["snapshot_no"] == 2
    assert snap2_id != snap1_id

    # Existen 2 instantáneas en base de datos
    snapshots = ChapterSnapshot.objects.filter(chapter=chapter).order_by("snapshot_no")
    assert snapshots.count() == 2

    # El snapshot 1 permanece inalterado
    s1 = snapshots[0]
    assert s1.snapshot_no == 1
    assert "Texto corregido y mejorado" not in s1.content_md
    assert "Borrador P05" in s1.content_md

    # El snapshot 2 tiene el nuevo texto y nuevo hash
    s2 = snapshots[1]
    assert s2.snapshot_no == 2
    assert "Texto corregido y mejorado para P05" in s2.content_md
    assert s2.payload_sha256 != s1.payload_sha256


@pytest.mark.django_db(transaction=True)
def test_concurrent_submit_creates_one_snapshot_and_returns_same_identity(snapshot_context):
    owner, _, _, project = snapshot_context
    outputs = _p01_to_p05_outputs()
    for idx, code in enumerate(["P01", "P02", "P03", "P04", "P05"]):
        save_step(owner, project.id, code, f"Contenido concurrente {code}", outputs[idx], idx, str(uuid4()))

    barrier = Barrier(2)

    def concurrent_submit():
        close_old_connections()
        try:
            thread_owner = get_user_model().objects.get(pk=owner.pk)
            barrier.wait(timeout=15)
            return submit_chapter(thread_owner, project.id, 1)["snapshot"]["id"]
        finally:
            close_old_connections()

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: concurrent_submit(), range(2)))

    assert results[0] == results[1]
    assert ChapterSnapshot.objects.filter(chapter__project=project, chapter__number=1).count() == 1


# ==============================================================================
# 4. Consulta Autorizada y Aislamiento por Membresía
# ==============================================================================


def test_snapshot_api_listing_and_detail_with_permissions(snapshot_context):
    owner, advisor, outsider, project = snapshot_context
    outputs = _p01_to_p05_outputs()

    for idx, code in enumerate(["P01", "P02", "P03", "P04", "P05"]):
        save_step(owner, project.id, code, f"Texto {code}", outputs[idx], idx, str(uuid4()))

    submit_chapter(owner, project.id, 1)

    # 1. El titular lista snapshots
    snaps_owner = list_chapter_snapshots(owner, project.id, 1)
    assert len(snaps_owner) == 1
    assert snaps_owner[0]["snapshot_no"] == 1

    # 2. El asesor asignado lista snapshots
    snaps_adv = list_chapter_snapshots(advisor, project.id, 1)
    assert len(snaps_adv) == 1

    # 3. Usuario ajeno NO puede listar
    with pytest.raises(DomainError) as exc_stranger:
        list_chapter_snapshots(outsider, project.id, 1)
    assert exc_stranger.value.code == "NOT_FOUND"

    snap_id = snaps_owner[0]["id"]

    # 4. Detalle de snapshot
    detail = get_chapter_snapshot(owner, project.id, 1, snap_id)
    assert detail["id"] == snap_id
    assert detail["chapter_no"] == 1
    assert detail["snapshot_no"] == 1
    assert "Texto P01" in detail["content_md"]
    assert len(detail["sections_manifest"]) == 5

    # El asesor comenta por snapshot y puede filtrar el historial por esa versión.
    client = Client()
    client.force_login(advisor)
    selected_text = detail["content_md"][:8]
    response_comment = client.post(
        f"/api/v1/projects/{project.id}/chapters/1/comments",
        data=json.dumps({
            "snapshot_id": snap_id,
            "start_cp": 0,
            "end_cp": len(selected_text),
            "selected_text": selected_text,
            "body": "Revisar la formulación de este fragmento.",
        }),
        content_type="application/json",
    )
    assert response_comment.status_code == 200, response_comment.content
    assert response_comment.json()["snapshot_id"] == snap_id
    response_comments = client.get(
        f"/api/v1/projects/{project.id}/chapters/1/comments?snapshot_id={snap_id}"
    )
    assert response_comments.status_code == 200
    assert len(response_comments.json()) == 1
    assert response_comments.json()[0]["snapshot_sha256"] == detail["payload_sha256"]

    chapter = Chapter.objects.get(project=project, number=1)
    revision = chapter.revisions.order_by("revision_no").first()
    comment_fields = {
        "author": advisor,
        "start_cp": 0,
        "end_cp": 1,
        "selected_text": "x",
        "body": "Comentario para probar la restricción XOR.",
    }
    with pytest.raises(IntegrityError), transaction.atomic():
        ReviewComment.objects.create(**comment_fields)
    with pytest.raises(IntegrityError), transaction.atomic():
        ReviewComment.objects.create(**comment_fields, revision=revision, snapshot=chapter.submitted_snapshot)

    # 5. Usuario ajeno NO puede consultar detalle
    with pytest.raises(DomainError) as exc_detail:
        get_chapter_snapshot(outsider, project.id, 1, snap_id)
    assert exc_detail.value.code == "NOT_FOUND"

    # 6. Comprobación vía cliente REST Django Ninja
    client = Client()
    client.force_login(owner)

    resp_list = client.get(f"/api/v1/projects/{project.id}/chapters/1/snapshots")
    assert resp_list.status_code == 200
    assert len(resp_list.json()) == 1

    resp_detail = client.get(f"/api/v1/projects/{project.id}/chapters/1/snapshots/{snap_id}")
    assert resp_detail.status_code == 200
    assert resp_detail.json()["snapshot_no"] == 1

    client_outsider = Client()
    client_outsider.force_login(outsider)
    resp_forbidden = client_outsider.get(f"/api/v1/projects/{project.id}/chapters/1/snapshots")
    assert resp_forbidden.status_code in (403, 404)


# ==============================================================================
# 5. Inmutabilidad en PostgreSQL: Bloqueo de UPDATE y DELETE por Trigger
# ==============================================================================


def test_postgresql_append_only_trigger_on_chapter_snapshot(snapshot_context):
    owner, _, _, project = snapshot_context
    outputs = _p01_to_p05_outputs()

    for idx, code in enumerate(["P01", "P02", "P03", "P04", "P05"]):
        save_step(owner, project.id, code, f"Texto {code}", outputs[idx], idx, str(uuid4()))

    submit_chapter(owner, project.id, 1)
    snapshot = ChapterSnapshot.objects.filter(chapter__project=project).first()
    assert snapshot is not None

    # Intentar modificar el snapshot directamente con SQL UPDATE debe fallar por el trigger
    with pytest.raises((ProgrammingError, DatabaseError)) as exc_update:
        with transaction.atomic():
            with connection.cursor() as cursor:
                cursor.execute(
                    "UPDATE chapter_snapshot SET content_md = 'Intento de alteración fraudulenta' WHERE id = %s",
                    [snapshot.id],
                )
    assert "Academic history is append-only" in str(exc_update.value)

    # Intentar eliminar el snapshot directamente con SQL DELETE debe fallar por el trigger
    with pytest.raises((ProgrammingError, DatabaseError)) as exc_delete:
        with transaction.atomic():
            with connection.cursor() as cursor:
                cursor.execute(
                    "DELETE FROM chapter_snapshot WHERE id = %s",
                    [snapshot.id],
                )
    assert "Academic history is append-only" in str(exc_delete.value)


def test_postgresql_append_only_trigger_on_chapter_document_revision(snapshot_context):
    owner, _, _, project = snapshot_context
    tree = {
        "schema_version": 1,
        "blocks": [{"type": "paragraph", "id": str(uuid4()), "text": "Documento inmutable."}],
    }
    saved = save_chapter_document(owner, project.id, 1, tree, expected_version=0, idempotency_key=str(uuid4()))
    revision = ChapterDocumentRevision.objects.get(pk=saved["revision_id"])

    with pytest.raises((ProgrammingError, DatabaseError)) as exc_update:
        with transaction.atomic():
            with connection.cursor() as cursor:
                cursor.execute(
                    "UPDATE chapter_document_revision SET content_md = 'alterado' WHERE id = %s",
                    [revision.id],
                )
    assert "Academic history is append-only" in str(exc_update.value)

    with pytest.raises((ProgrammingError, DatabaseError)) as exc_delete:
        with transaction.atomic():
            with connection.cursor() as cursor:
                cursor.execute(
                    "DELETE FROM chapter_document_revision WHERE id = %s",
                    [revision.id],
                )
    assert "Academic history is append-only" in str(exc_delete.value)


def test_snapshot_freezes_citation_manifest(snapshot_context):
    import hashlib
    import json

    owner, advisor, _, project = snapshot_context
    outputs = _p01_to_p05_outputs()
    for idx, code in enumerate(["P01", "P02", "P03", "P04", "P05"]):
        save_step(owner, project.id, code, f"Texto {code}", outputs[idx], idx, str(uuid4()))
    source = SourceDocument.objects.create(
        project=project, kind=SourceDocument.Kind.EXTERNAL,
        record_provider=SourceDocument.RecordProvider.ALICIA, provider_record_id="alicia-snapshot-citation",
        title="Fuente congelada en snapshot", authors_json=["Autora"], doi="10.5678/snapshot-citation",
    )
    block_id = str(uuid4())
    tree = {"schema_version": 1, "blocks": [{
        "type": "paragraph", "id": block_id, "text": "Una afirmación con cita."}]}
    saved = save_chapter_document(
        owner, project.id, 1, tree, 0, str(uuid4()),
        citations=[{"group_id": str(uuid4()), "block_id": block_id, "position": 17,
                    "sources": [{"source_id": str(source.id), "locator": "p. 12"}]}],
    )
    snapshot_result = submit_chapter(owner, project.id, 1)
    snapshot = Chapter.objects.get(project=project, number=1).submitted_snapshot
    detail = get_chapter_snapshot(owner, project.id, 1, snapshot.id)
    assert detail["document_revision_id"] == saved["revision_id"]
    assert detail["citations"][0]["sources"][0]["reference"]["title"] == "Fuente congelada en snapshot"
    assert snapshot.payload_sha256 == snapshot_result["snapshot"]["payload_sha256"]
    expected_payload = {
        "chapter_no": snapshot.chapter.number,
        "content_md": snapshot.content_md,
        "format_version": snapshot.format_version,
        "project_title": snapshot.project_title,
        "sections_manifest": snapshot.sections_manifest,
        "state_version": snapshot.state_version,
        "workflow_id": str(snapshot.workflow_id) if snapshot.workflow_id else None,
        "workflow_content_sha256": snapshot.workflow.content_sha256 if snapshot.workflow_id else "",
        "workflow_version_no": snapshot.workflow_version_no,
        "document_revision_id": str(snapshot.document_revision_id),
        "document_revision_sha256": snapshot.document_revision.content_sha256,
        "citation_manifest_sha256": snapshot.document_revision.citation_manifest_sha256,
        "document_artifacts_manifest": [],
    }
    expected_hash = hashlib.sha256(json.dumps(
        expected_payload, sort_keys=True, ensure_ascii=False,
    ).encode("utf-8")).hexdigest()
    assert snapshot.payload_sha256 == expected_hash

    decide_chapter(advisor, project.id, 1, "APPROVE")
    exported = Document(export_approved_snapshot_docx(owner, project.id, 1))
    exported_text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
    assert "[1]" in exported_text
    assert "Referencias" in exported_text
    assert "Fuente congelada en snapshot" in exported_text


def test_snapshot_and_approved_export_freeze_table_and_figure_versions(snapshot_context, tmp_path):
    owner, advisor, _, project = snapshot_context
    outputs = _p01_to_p05_outputs()
    for idx, code in enumerate(["P01", "P02", "P03", "P04", "P05"]):
        save_step(owner, project.id, code, f"Texto {code}", outputs[idx], idx, str(uuid4()))
    with override_settings(MEDIA_ROOT=tmp_path):
        dataset = create_dataset(
            owner, project.id, "Datos para exportación",
            SimpleUploadedFile("datos.csv", b"grupo,valor\nA,2\nB,4\n", content_type="text/csv"), uuid4(),
        )
        table = create_table_recipe(owner, project.id, dataset["dataset"]["id"], 1, "Tabla de resultados",
                                    {"columns": ["grupo", "valor"], "limit": 10}, uuid4())
        figure = create_figure(
            owner, project.id, dataset["dataset"]["id"], table["recipe"]["id"], 1, "Barras", {
                "chart_type": "BAR", "x_column": "grupo", "y_column": "valor", "title": "Valores por grupo",
                "x_label": "Grupo", "y_label": "Valor", "alt_text": "A tiene 2 y B tiene 4.",
            }, uuid4(),
        )
        table_block, figure_block = str(uuid4()), str(uuid4())
        tree = {"schema_version": 2, "blocks": [
            {"type": "heading", "id": str(uuid4()), "level": 2, "text": "Resultados"},
            {"type": "data_table", "id": table_block,
             "table_version_id": table["version"]["id"],
             "result_sha256": table["version"]["result_sha256"], "caption": "Tabla 1. Valores registrados."},
            {"type": "figure", "id": figure_block,
             "figure_version_id": figure["version"]["id"],
             "output_sha256": figure["version"]["output_sha256"], "caption": "Figura 1. Valores por grupo."},
        ]}
        save_chapter_document(owner, project.id, 1, tree, 0, str(uuid4()))
        submitted = submit_chapter(owner, project.id, 1)
        snapshot = Chapter.objects.get(project=project, number=1).submitted_snapshot
        assert snapshot.format_version == 3
        assert {item["block_id"] for item in snapshot.document_artifacts_manifest} == {table_block, figure_block}
        assert snapshot.payload_sha256 == submitted["snapshot"]["payload_sha256"]
        assert ResearchFigureVersion.objects.filter(figure_id=figure["figure"]["id"]).count() == 1

        # Versiones posteriores de tabla y figura no cambian el snapshot que guardó v1.
        revised_table = create_table_recipe(
            owner, project.id, dataset["dataset"]["id"], 1, "Tabla de resultados",
            {"columns": ["grupo"], "limit": 1}, uuid4(), expected_version=1,
        )
        assert revised_table["version"]["version_no"] == 2
        revised_figure = create_figure(
            owner, project.id, dataset["dataset"]["id"], table["recipe"]["id"], 1, "Barras", {
                "chart_type": "BAR", "x_column": "grupo", "y_column": "valor", "title": "Nuevo título",
                "x_label": "Grupo", "y_label": "Valor", "alt_text": "Versión posterior.",
            }, uuid4(), expected_version=1,
        )
        assert revised_figure["version"]["version_no"] == 2
        decide_chapter(advisor, project.id, 1, "APPROVE")
        exported = Document(export_approved_snapshot_docx(owner, project.id, 1))
        exported_text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
        assert "Tabla 1. Valores registrados." in exported_text
        assert "Figura 1. Valores por grupo." in exported_text
        assert len(exported.tables) == 1
        assert [cell.text for cell in exported.tables[0].rows[1].cells] == ["A", "2"]
        assert len(exported.inline_shapes) == 1
        assert exported.inline_shapes[0]._inline.docPr.get("descr") == "A tiene 2 y B tiene 4."
