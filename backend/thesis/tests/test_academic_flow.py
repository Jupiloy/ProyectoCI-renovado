from uuid import uuid4
from unittest.mock import patch
import json
import asyncio
import hashlib

import pytest
from django.contrib.auth import get_user_model
from django.test import Client
from docx import Document

from thesis.models import Chapter, ChapterRevision, ProjectParticipant, ReviewComment, ReviewEvent, SourceDocument, SourceExcerpt
from thesis.services.errors import DomainError
from thesis.services.evidence import check_source_doi, locate_quote
from thesis.services.evidence_links import create_evidence_link, review_evidence_link
from thesis.services.exports import export_approved_snapshot_docx, export_draft_docx
from thesis.services.chapter_documents import get_chapter_document, save_chapter_document
from thesis.services.projects import create_project
from thesis.services.reviews import add_comment, assign_advisor, decide_chapter, list_comments, resolve_comment
from thesis.services.steps import get_chapter, save_step, submit_chapter
from thesis.services.step_contracts import P20Output
from pydantic import ValidationError


@pytest.fixture
def context(db):
    users = get_user_model().objects
    owner = users.create_user(email="flow-owner@example.org", password="pass", display_name="Ana")
    advisor = users.create_user(email="flow-advisor@example.org", password="pass", display_name="Asesora")
    outsider = users.create_user(email="flow-outsider@example.org", password="pass", display_name="Otra")
    project = create_project(owner, "Tesis")
    return owner, advisor, outsider, project


def outputs():
    candidates = [{"id": i, "titulo_tentativo": f"Tema candidato número {i}", "pros": ["Acceso"],
                   "contras": ["Tiempo"], "viabilidad_preliminar": "MEDIA"} for i in range(1, 6)]
    score = {"acceso_a_datos": 4, "relevancia_profesional": 4, "brecha_conocimiento": 4,
             "recursos_disponibles": 4}
    matrix = [{"tema_id": i, "titulo": f"Tema candidato número {i}", "puntajes": score,
               "puntaje_total": 16, "justificacion_eleccion": "Se dispone de datos y tiempo suficiente."}
              for i in range(1, 4)]
    introduction = "\n\n".join(["palabra " * 90] * 3).strip()
    return [
        {"temas_propuestos": candidates, "tema_preseleccionado_id": 1},
        {"matriz_comparativa": matrix, "tema_ganador_id": 1, "tema_ganador_titulo": matrix[0]["titulo"]},
        {"ods_principal_numero": 4, "ods_principal_nombre": "Educación de calidad",
         "meta_ods_especifica": "4.4", "argumento_alineacion": "El trabajo aborda acceso a formación y necesidades de una comunidad educativa concreta.",
         "linea_utp_vinculada": "Tecnología educativa"},
        {"titulo_oficial": "Uso de herramientas digitales en estudiantes de Piura durante 2026",
         "conteo_palabras": 10, "variables_confirmadas": ["herramientas digitales"],
         "unidad_analisis_confirmada": "estudiantes", "lugar_confirmado": "Piura", "periodo_confirmado": 2026},
        {"contenido_markdown": introduction, "num_parrafos": 3, "palabras_totales": 270},
    ]


def test_draft_export_contains_saved_revision_and_enforces_membership(context):
    owner, _, outsider, project = context
    save_step(owner, project.id, "P01", "# Tema propio\nTexto redactado por el tesista.", outputs()[0], 0, str(uuid4()))
    document = Document(export_draft_docx(owner, project.id))
    assert round(document.sections[0].left_margin.cm, 1) == 3.0
    assert document.styles["Normal"].font.name == "Arial"
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Texto redactado por el tesista." in body
    assert "Capítulo 1" in body
    assert "Capítulo 2" not in body
    client = Client()
    client.force_login(owner)
    response = client.get(f"/api/v1/projects/{project.id}/export/draft.docx")
    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    with pytest.raises(DomainError) as denied:
        export_draft_docx(outsider, project.id)
    assert denied.value.code == "NOT_FOUND"
    client.force_login(outsider)
    assert client.get(f"/api/v1/projects/{project.id}/export/draft.docx").status_code == 404


def test_approved_snapshot_export_is_frozen_and_requires_approved_snapshot(context):
    owner, advisor, outsider, project = context
    assign_advisor(owner, project.id, advisor.email)
    for index, output in enumerate(outputs(), start=1):
        save_step(owner, project.id, f"P{index:02d}", f"Texto congelado de P{index:02d}", output,
                  index - 1, str(uuid4()))
    submitted = submit_chapter(owner, project.id, 1)

    with pytest.raises(DomainError) as pending:
        export_approved_snapshot_docx(owner, project.id, 1)
    assert pending.value.code == "VERSION_CONFLICT"

    decide_chapter(advisor, project.id, 1, "APPROVE")
    stream = export_approved_snapshot_docx(owner, project.id, 1)
    document = Document(stream)
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Capítulo 1" in body
    assert "Texto congelado de P01" in body
    assert "Texto congelado de P05" in body
    assert submitted["snapshot"]["id"] in body
    assert submitted["snapshot"]["payload_sha256"] in body
    assert "temas_propuestos" not in body
    assert round(document.sections[0].left_margin.cm, 1) == 3.0

    client = Client()
    client.force_login(owner)
    response = client.get(f"/api/v1/projects/{project.id}/chapters/1/export/approved.docx")
    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert response["Content-Disposition"].endswith('filename="capitulo-1-aprobado.docx"')

    with pytest.raises(DomainError) as denied:
        export_approved_snapshot_docx(outsider, project.id, 1)
    assert denied.value.code == "NOT_FOUND"
    client.force_login(outsider)
    assert client.get(f"/api/v1/projects/{project.id}/chapters/1/export/approved.docx").status_code == 404

    decide_chapter(advisor, project.id, 1, "REOPEN", "Corrección posterior")
    with pytest.raises(DomainError) as reopened:
        export_approved_snapshot_docx(owner, project.id, 1)
    assert reopened.value.code == "VERSION_CONFLICT"


def test_canonical_chapter_document_versioning_snapshot_and_export(context):
    owner, advisor, outsider, project = context
    assign_advisor(owner, project.id, advisor.email)
    for index, output in enumerate(outputs(), start=1):
        save_step(owner, project.id, f"P{index:02d}", f"Texto legado de P{index:02d}", output,
                  index - 1, str(uuid4()))

    client = Client()
    client.force_login(owner)
    legacy = client.get(f"/api/v1/projects/{project.id}/chapters/1/document")
    assert legacy.status_code == 200
    assert legacy.json()["is_legacy_projection"] is True
    assert "P01. Mapa de temas viables" in legacy.json()["content_md"]

    first_tree = {
        "schema_version": 1,
        "blocks": [
            {"type": "heading", "id": str(uuid4()), "level": 1, "text": "Capítulo canónico"},
            {"type": "paragraph", "id": str(uuid4()), "text": "Texto editable que debe quedar congelado."},
            {"type": "bullet_list", "id": str(uuid4()), "items": ["Elemento uno", "Elemento dos"]},
        ],
    }
    idempotency_key = str(uuid4())
    save_response = client.put(
        f"/api/v1/projects/{project.id}/chapters/1/document",
        data=json.dumps({"content_tree": first_tree, "expected_version": 0,
                         "idempotency_key": idempotency_key}),
        content_type="application/json",
    )
    assert save_response.status_code == 200
    first_saved = save_response.json()
    assert first_saved["document_version"] == 1
    assert first_saved["content_tree"] == first_tree
    assert "# Capítulo canónico" in first_saved["content_md"]
    repeated = client.put(
        f"/api/v1/projects/{project.id}/chapters/1/document",
        data=json.dumps({"content_tree": first_tree, "expected_version": 0,
                         "idempotency_key": idempotency_key}),
        content_type="application/json",
    )
    assert repeated.status_code == 200
    assert repeated.json()["idempotent"] is True
    assert repeated.json()["revision_id"] == first_saved["revision_id"]

    current_profile_version = first_saved["state_version"]
    save_step(owner, project.id, "P05", "Revisión de paso posterior", outputs()[4],
              current_profile_version, str(uuid4()))
    with pytest.raises(DomainError) as stale_document:
        submit_chapter(owner, project.id, 1)
    assert stale_document.value.code == "CONTEXT_STALE"

    second_tree = {
        "schema_version": 1,
        "blocks": [
            {"type": "heading", "id": first_tree["blocks"][0]["id"], "level": 1, "text": "Capítulo canónico"},
            {"type": "paragraph", "id": first_tree["blocks"][1]["id"], "text": "Texto editado y revisado."},
        ],
    }
    second_saved = save_chapter_document(
        owner, project.id, 1, second_tree, expected_version=1, idempotency_key=str(uuid4()),
    )
    submitted = submit_chapter(owner, project.id, 1)
    snapshot = Chapter.objects.get(project=project, number=1).submitted_snapshot
    assert snapshot.format_version == 2
    assert str(snapshot.document_revision_id) == second_saved["revision_id"]
    assert snapshot.content_md == second_saved["content_md"]
    assert all(item["rendered_in_text"] is False for item in snapshot.sections_manifest)
    assert submitted["snapshot"]["document_revision_id"] == second_saved["revision_id"]

    decide_chapter(advisor, project.id, 1, "APPROVE")
    approved_doc = Document(export_approved_snapshot_docx(owner, project.id, 1))
    approved_body = "\n".join(paragraph.text for paragraph in approved_doc.paragraphs)
    assert "Capítulo canónico" in approved_body
    assert "Texto editado y revisado." in approved_body
    assert "Texto editable que debe quedar congelado." not in approved_body
    assert "##" not in approved_body

    outsider_client = Client()
    outsider_client.force_login(outsider)
    assert outsider_client.get(f"/api/v1/projects/{project.id}/chapters/1/document").status_code == 404


def test_p20_accepts_ieee_and_preserves_legacy_apa():
    references = ["Referencia de prueba"]
    for style in ("IEEE", "APA_7"):
        assert P20Output.model_validate({
            "estilo_citacion": style, "referencias_formateadas": references,
        }).estilo_citacion == style
    with pytest.raises(ValidationError):
        P20Output.model_validate({"estilo_citacion": "OTRO", "referencias_formateadas": references})


def test_full_chapter_review_with_immutable_revisions(context):
    owner, advisor, outsider, project = context
    assign_advisor(owner, project.id, advisor.email)
    assert ProjectParticipant.objects.filter(project=project, user=advisor, active=True).exists()
    with pytest.raises(DomainError) as denied:
        assign_advisor(outsider, project.id, advisor.email)
    assert denied.value.code == "NOT_FOUND"

    for index, output in enumerate(outputs(), start=1):
        code = f"P{index:02d}"
        result = save_step(owner, project.id, code, f"Texto de capítulo en paso {code}", output,
                           index - 1, str(uuid4()))
        assert result["revision_no"] == index
    chapter = get_chapter(owner, project.id, 1)
    assert chapter["latest_revision"]["step_code"] == "P05"
    submitted = submit_chapter(owner, project.id, 1)
    assert submitted["status"] == "READY"
    snapshot1_id = submitted["snapshot"]["id"]
    with pytest.raises(DomainError) as blocked:
        decide_chapter(advisor, project.id, 1, "REQUEST_CHANGES")
    assert blocked.value.code == "VALIDATION_ERROR"
    revision = ChapterRevision.objects.get(chapter__project=project, chapter__number=1, revision_no=5)
    snapshot1 = submitted["snapshot"]
    comment = add_comment(
        advisor,
        project.id,
        1,
        0,
        5,
        "Texto",
        "Aclare la formulación del capítulo",
        snapshot_id=snapshot1_id,
    )
    assert comment["revision_id"] is None
    assert comment["snapshot_id"] == snapshot1_id
    assert comment["snapshot_sha256"] == snapshot1["payload_sha256"]
    assert str(ReviewComment.objects.get(pk=comment["comment_id"]).snapshot_id) == snapshot1_id
    with pytest.raises(DomainError) as stale_snapshot:
        add_comment(advisor, project.id, 1, 0, 5, "Texto", "Comentario obsoleto", snapshot_id=uuid4())
    assert stale_snapshot.value.code == "VERSION_CONFLICT"
    with pytest.raises(DomainError) as denied_comment:
        add_comment(outsider, project.id, 1, 0, 5, "Texto", "Sin permiso", snapshot_id=snapshot1_id)
    assert denied_comment.value.code == "NOT_FOUND"
    with pytest.raises(DomainError) as incorrect_selection:
        add_comment(advisor, project.id, 1, 0, 5, "texto", "Texto no coincidente", snapshot_id=snapshot1_id)
    assert incorrect_selection.value.code == "VALIDATION_ERROR"
    with pytest.raises(DomainError) as blocked:
        decide_chapter(advisor, project.id, 1, "APPROVE")
    assert blocked.value.code == "ADVISOR_UNRESOLVED_COMMENTS"
    assert decide_chapter(advisor, project.id, 1, "REQUEST_CHANGES")["status"] == "CHANGES_REQUESTED"
    request_event = ReviewEvent.objects.get(chapter__project=project, action=ReviewEvent.Action.REQUEST_CHANGES)
    assert str(request_event.snapshot_id) == snapshot1_id
    assert request_event.revision_id is None
    another = save_step(owner, project.id, "P05", "Texto corregido del capítulo", outputs()[4],
                        5, str(uuid4()))
    assert another["revision_no"] == 6
    assert get_chapter(owner, project.id, 1)["status"] == "DRAFT"
    submitted2 = submit_chapter(owner, project.id, 1)
    snapshot2_id = submitted2["snapshot"]["id"]
    assert len(list_comments(advisor, project.id, 1, snapshot_id=snapshot1_id)) == 1
    assert list_comments(advisor, project.id, 1, snapshot_id=snapshot2_id) == []
    assert decide_chapter(advisor, project.id, 1, "APPROVE")["status"] == "APPROVED"
    approved_chapter = Chapter.objects.get(project=project, number=1)
    assert str(approved_chapter.approved_snapshot_id) == snapshot2_id
    assert approved_chapter.approved_revision_id is None
    approved_event = ReviewEvent.objects.get(chapter=approved_chapter, action=ReviewEvent.Action.APPROVE)
    assert str(approved_event.snapshot_id) == snapshot2_id
    assert approved_event.revision_id is None
    assert decide_chapter(advisor, project.id, 1, "REOPEN", "Revisar una corrección adicional")["status"] == "DRAFT"
    reopen_event = ReviewEvent.objects.get(chapter=approved_chapter, action=ReviewEvent.Action.REOPEN)
    assert str(reopen_event.snapshot_id) == snapshot2_id
    assert Chapter.objects.get(pk=approved_chapter.pk).approved_snapshot_id is None
    assert ReviewEvent.objects.filter(chapter__project=project).count() == 5
    assert resolve_comment(advisor, project.id, comment["comment_id"])["status"] == "RESOLVED"


def test_legacy_ready_chapter_without_snapshot_keeps_revision_review_path(context):
    owner, advisor, _, project = context
    assign_advisor(owner, project.id, advisor.email)
    saved = save_step(owner, project.id, "P01", "Texto histórico revisable", outputs()[0], 0, str(uuid4()))
    chapter = Chapter.objects.get(project=project, number=1)
    chapter.submit()
    chapter.save(update_fields=["status", "updated_at"])

    comment = add_comment(advisor, project.id, 1, 0, 5, "Texto", "Aclarar el sentido histórico")
    assert comment["revision_id"] == saved["revision_id"]
    assert comment["snapshot_id"] is None
    assert list_comments(advisor, project.id, 1)[0]["revision_id"] == saved["revision_id"]
    resolve_comment(advisor, project.id, comment["comment_id"])

    decide_chapter(advisor, project.id, 1, "APPROVE")
    chapter.refresh_from_db()
    assert chapter.approved_revision_id == saved["revision_id"]
    assert chapter.approved_snapshot_id is None
    approval = ReviewEvent.objects.get(chapter=chapter, action=ReviewEvent.Action.APPROVE)
    assert approval.revision_id == saved["revision_id"]
    assert approval.snapshot_id is None
    with pytest.raises(DomainError) as no_snapshot_export:
        export_approved_snapshot_docx(owner, project.id, 1)
    assert no_snapshot_export.value.code == "VERSION_CONFLICT"


def test_contract_and_idempotency(context):
    owner, _, _, project = context
    invalid = outputs()[0]
    invalid["tema_preseleccionado_id"] = 9
    with pytest.raises(DomainError) as error:
        save_step(owner, project.id, "P01", "Texto", invalid, 0, str(uuid4()))
    assert error.value.code == "VALIDATION_SCHEMA_ERROR"
    key = str(uuid4())
    valid = outputs()[0]
    first = save_step(owner, project.id, "P01", "Texto", valid, 0, key)
    second = save_step(owner, project.id, "P01", "Texto", valid, 0, key)
    assert first["revision_id"] == second["revision_id"]
    assert second["idempotent"] is True
    with pytest.raises(DomainError) as error:
        save_step(owner, project.id, "P03", "Texto", outputs()[2], 1, str(uuid4()))
    assert error.value.code == "PREREQUISITE_LOCKED"


def test_registry_and_quote_report_only_candidates(context):
    owner, advisor, _, project = context
    source = SourceDocument.objects.create(
        project=project, kind=SourceDocument.Kind.UPLOAD, title="Etiqueta del usuario",
        authors_json=[], storage_key="sample.pdf", sha256="a" * 64,
        doi="10.1234/example", verification_status=SourceDocument.Status.TEXT_CHECKED,
    )
    excerpt_text = "Esta es la frase exacta del documento."
    excerpt = SourceExcerpt.objects.create(source=source, chunk_no=1, page_no=4,
                                           text=excerpt_text, text_sha256=hashlib.sha256(excerpt_text.encode()).hexdigest())
    with patch("thesis.services.evidence.get_doi_record", return_value={
        "DOI": "10.1234/example", "title": ["Título registrado"],
        "author": [{"given": "Ana", "family": "Prueba"}],
    }):
        checked = check_source_doi(owner, project.id, source.id)
    assert checked["registered_title"] == "Título registrado"
    assert checked["title_similarity_candidate"] is None
    assert checked["cover_metadata_verified"] is False
    found = locate_quote(owner, project.id, source.id, "frase exacta del documento")
    assert found["matches"][0]["physical_page"] == 4
    assert found["matches"][0]["match_type"] == "EXACT"
    assert found["scientific_claims_verified"] is False
    with pytest.raises(DomainError):
        locate_quote(advisor, project.id, source.id, "frase exacta del documento")
    assign_advisor(owner, project.id, advisor.email)
    match = found["matches"][0]
    link = create_evidence_link(owner, project.id, excerpt.id,
                                "La frase exacta describe el documento", match["start_char"],
                                match["end_char"], match["matched_text"])
    assert link["status"] == "CANDIDATE"
    assert link["physical_page"] == 4
    with pytest.raises(DomainError):
        review_evidence_link(owner, project.id, link["evidence_link_id"], "CONFIRMED")
    reviewed = review_evidence_link(advisor, project.id, link["evidence_link_id"], "CONFIRMED")
    assert reviewed["status"] == "CONFIRMED"
    assert reviewed["automatically_verified"] is False


def test_api_and_mcp_expose_new_workflow(context):
    owner, advisor, _, project = context
    client = Client()
    client.force_login(owner)
    assigned = client.post(f"/api/v1/projects/{project.id}/advisor",
                           data=json.dumps({"email": advisor.email}), content_type="application/json")
    assert assigned.status_code == 200, assigned.content
    saved = client.put(f"/api/v1/projects/{project.id}/steps/P01", data=json.dumps({
        "content_md": "Texto confirmado por Ana", "output_data": outputs()[0],
        "expected_version": 0, "idempotency_key": str(uuid4()),
    }), content_type="application/json")
    assert saved.status_code == 200, saved.content
    client.force_login(advisor)
    dashboard = client.get("/api/v1/projects/advisor/dashboard")
    assert dashboard.status_code == 200, dashboard.content
    assert dashboard.json()[0]["project_id"] == str(project.id)
    chapter = client.get(f"/api/v1/projects/{project.id}/chapters/1")
    assert chapter.json()["latest_revision"]["step_code"] == "P01"

    from mcp_server.server import app
    tool_names = {tool.name for tool in asyncio.run(app.list_tools())}
    assert {"guardar_paso", "consultar_capitulo", "decidir_capitulo", "localizar_frase", "cotejar_doi"} <= tool_names
