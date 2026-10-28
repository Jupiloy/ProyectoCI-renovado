import hashlib
from io import BytesIO
from unittest.mock import patch

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, override_settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from thesis.integrations.docx import inspect_docx
from thesis.models import ManuscriptDocument
from thesis.services.errors import DomainError
from thesis.services.manuscripts import get_manuscript, list_manuscript_citations, read_manuscript_section, upload_docx_pending
from thesis.services.projects import create_project
from thesis.services.document_jobs import run_next_document_job


def sample_docx() -> bytes:
    document = Document()
    document.add_paragraph("Esta es una idea propia (Autor, 2020), escrita solo como texto.")
    paragraph = document.add_paragraph("Cita estructurada: ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "CITATION Smith2021 \\p 15")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "(Smith, 2021, p. 15)"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    zotero = OxmlElement("w:fldSimple")
    zotero.set(qn("w:instr"), 'ADDIN ZOTERO_ITEM CSL_CITATION {"citationItems":[{"id":42,"locator":"7","label":"page"}]}')
    run2 = OxmlElement("w:r")
    text2 = OxmlElement("w:t")
    text2.text = "(Doe, 2022, p. 7)"
    run2.append(text2)
    zotero.append(run2)
    paragraph._p.append(zotero)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture
def context(db):
    user = get_user_model().objects.create_user(email="docx@example.org", password="test-pass", display_name="Ana")
    other = get_user_model().objects.create_user(email="other-docx@example.org", password="test-pass", display_name="Otra")
    return user, other, create_project(user, "Tesis DOCX")


def test_docx_fields_and_conversion_are_kept_separate(context, tmp_path):
    user, other, project = context
    content = sample_docx()
    with override_settings(MEDIA_ROOT=tmp_path):
        result = upload_docx_pending(user, project.id, "Capítulo de Ana", SimpleUploadedFile("chapter.docx", content))
        assert result["upload_status"] == "DOCX_QUEUED"
        assert run_next_document_job() is True
        job = ManuscriptDocument.objects.get(pk=result["manuscript_id"]).document_jobs.get()
        assert job.status == "SUCCEEDED", job.error_message
        detail = get_manuscript(user, project.id, result["manuscript_id"])
        stored = tmp_path / ManuscriptDocument.objects.get(pk=result["manuscript_id"]).storage_key
        assert stored.read_bytes() == content
        assert result["sha256"] == hashlib.sha256(content).hexdigest()
        duplicate = upload_docx_pending(user, project.id, "Repetido", SimpleUploadedFile("chapter.docx", content))
    assert duplicate["upload_status"] == "DUPLICATE_DOCX"
    assert detail["citation_count"] == 2
    assert "idea propia" in detail["markdown"]
    assert {citation["kind"] for citation in detail["citations"]} == {"WORD", "ZOTERO"}
    assert detail["citations"][1]["citation_items"][0]["locator"] == "7"
    assert all(citation["verified_against_source"] is False for citation in detail["citations"])
    section = read_manuscript_section(user, project.id, result["manuscript_id"], offset=0, limit=12)
    assert section["markdown"] == detail["markdown"][:12]
    assert section["next_offset"] == 12
    citations_page = list_manuscript_citations(user, project.id, result["manuscript_id"], limit=1)
    assert citations_page["total_citations"] == 2
    assert len(citations_page["citations"]) == 1
    with pytest.raises(DomainError) as denied:
        get_manuscript(other, project.id, result["manuscript_id"])
    assert denied.value.code == "NOT_FOUND"


def test_docx_api_and_invalid_archive(context, tmp_path):
    user, _, project = context
    client = Client(enforce_csrf_checks=True)
    client.force_login(user)
    client.get("/api/v1/docs")
    csrf = client.cookies["csrftoken"].value
    with override_settings(MEDIA_ROOT=tmp_path):
        response = client.post(
            f"/api/v1/projects/{project.id}/manuscripts/upload",
            {"label": "Borrador", "file": SimpleUploadedFile("chapter.docx", sample_docx())},
            HTTP_X_CSRFTOKEN=csrf,
        )
        assert response.status_code == 202, response.content
        assert response.json()["job"]["status"] == "QUEUED"
        with patch("thesis.services.manuscripts.convert_docx", return_value="Borrador convertido"), \
             patch("thesis.services.manuscripts.inspect_docx", return_value=[]):
            assert run_next_document_job() is True
        assert client.get(f"/api/v1/projects/{project.id}/manuscripts").json()[0]["processing_status"] == "READY"
        invalid = client.post(
            f"/api/v1/projects/{project.id}/manuscripts/upload",
            {"label": "Falso", "file": SimpleUploadedFile("fake.docx", b"not an archive")},
            HTTP_X_CSRFTOKEN=csrf,
        )
    assert invalid.status_code == 422
    assert ManuscriptDocument.objects.filter(project=project).count() == 1
