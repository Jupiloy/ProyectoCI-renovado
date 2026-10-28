"""Tablas descriptivas reproducibles desde versiones inmutables de datasets."""

import csv
import hashlib
import json
import math
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from functools import cmp_to_key
from pathlib import Path
from uuid import UUID

from django.db import IntegrityError, transaction

from thesis.models import (
    ResearchDatasetVersion, ResearchTableRecipe, ResearchTableRecipeVersion,
)
from thesis.services.access import require_project
from thesis.services.errors import DomainError
from thesis.services.research_data import get_dataset_file


ENGINE_VERSION = "b8.2.1"
MAX_RESULT_ROWS = 10_000
MAX_RESULT_BYTES = 4 * 1024 * 1024
_FILTER_OPERATORS = {"EQ", "NE", "GT", "GTE", "LT", "LTE", "CONTAINS", "IS_EMPTY", "NOT_EMPTY"}
_SORT_DIRECTIONS = {"ASC", "DESC"}


def _canonical_json(value) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool) or value is None or isinstance(value, (str, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return value
    if isinstance(value, Decimal):
        if not value.is_finite():
            return None
        return format(value, "f")
    return str(value)


def _read_source_rows(path: Path, version: ResearchDatasetVersion) -> tuple[list[str], list[dict]]:
    headers = [column["name"] for column in version.columns_json]
    rows = []
    try:
        if version.file_format == ResearchDatasetVersion.Format.CSV:
            with path.open("r", encoding="utf-8-sig", newline="") as stream:
                reader = csv.reader(stream, strict=True)
                actual_headers = [str(value).strip() for value in next(reader, [])]
                if actual_headers != headers:
                    raise DomainError("OUTPUT_INVALID", "Los encabezados físicos no coinciden con el esquema registrado")
                for row in reader:
                    if not any(str(value).strip() for value in row):
                        continue
                    if len(row) > len(headers):
                        raise DomainError("OUTPUT_INVALID", "Una fila CSV excede el número de columnas registrado")
                    row.extend([""] * (len(headers) - len(row)))
                    rows.append(dict(zip(headers, row)))
        elif version.file_format == ResearchDatasetVersion.Format.XLSX:
            from openpyxl import load_workbook

            with zipfile.ZipFile(path) as archive:
                names = {item.filename for item in archive.infolist()}
                if "xl/workbook.xml" not in names:
                    raise DomainError("OUTPUT_INVALID", "El archivo XLSX no conserva su estructura registrada")
            values_workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            formulas_workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                sheet_name = version.import_metadata_json.get("selected_sheet")
                if sheet_name not in values_workbook.sheetnames or sheet_name not in formulas_workbook.sheetnames:
                    raise DomainError("OUTPUT_INVALID", "La hoja registrada ya no está disponible en el XLSX")
                sheet = values_workbook[sheet_name]
                formula_sheet = formulas_workbook[sheet_name]
                header_row = next(sheet.iter_rows(values_only=True), None)
                actual_headers = [str(value).strip() if value is not None else "" for value in (header_row or ())]
                while actual_headers and not actual_headers[-1]:
                    actual_headers.pop()
                if actual_headers != headers:
                    raise DomainError("OUTPUT_INVALID", "Los encabezados físicos no coinciden con el esquema registrado")
                value_rows = sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True)
                formula_rows = formula_sheet.iter_rows(min_row=2, max_col=len(headers))
                for value_row, formula_row in zip(value_rows, formula_rows):
                    values = list(value_row)
                    has_formula = any(cell.data_type == "f" for cell in formula_row)
                    if not any(value not in (None, "") for value in values) and not has_formula:
                        continue
                    values.extend([None] * (len(headers) - len(values)))
                    rows.append({name: _json_value(value) for name, value in zip(headers, values)})
            finally:
                values_workbook.close()
                formulas_workbook.close()
        else:
            raise DomainError("INVALID_DATASET", "Formato de dataset no admitido para recetas")
    except (UnicodeDecodeError, csv.Error, zipfile.BadZipFile, OSError) as exc:
        raise DomainError("OUTPUT_INVALID", "No se pudo reconstruir la tabla desde el archivo original") from exc
    if len(rows) != version.row_count:
        raise DomainError("OUTPUT_INVALID", "El número de filas reconstruido no coincide con la versión importada")
    return headers, rows


def _validate_spec(spec, headers: list[str], type_by_column: dict[str, str]) -> dict:
    if not isinstance(spec, dict) or set(spec) - {"columns", "filters", "sort", "limit"}:
        raise DomainError("VALIDATION_ERROR", "La receta solo admite columns, filters, sort y limit")
    columns = spec.get("columns")
    if not isinstance(columns, list) or not columns or len(columns) > len(headers):
        raise DomainError("VALIDATION_ERROR", "columns debe ser una lista no vacía de columnas de origen")
    if any(not isinstance(name, str) or name not in headers for name in columns) or len(set(columns)) != len(columns):
        raise DomainError("VALIDATION_ERROR", "columns contiene nombres inválidos o repetidos")
    limit = spec.get("limit")
    if type(limit) is not int or not 1 <= limit <= MAX_RESULT_ROWS:
        raise DomainError("VALIDATION_ERROR", f"limit es obligatorio y debe estar entre 1 y {MAX_RESULT_ROWS}")
    filters = spec.get("filters", [])
    if not isinstance(filters, list) or len(filters) > 64:
        raise DomainError("VALIDATION_ERROR", "filters debe contener como máximo 64 reglas")
    clean_filters = []
    for item in filters:
        if not isinstance(item, dict):
            raise DomainError("VALIDATION_ERROR", "Cada filtro debe ser un objeto")
        operator = item.get("operator")
        if not isinstance(operator, str):
            raise DomainError("VALIDATION_ERROR", "El operador de filtro debe ser texto")
        allowed = {"column", "operator"} if operator in {"IS_EMPTY", "NOT_EMPTY"} else {"column", "operator", "value"}
        if set(item) != allowed or item.get("column") not in headers or operator not in _FILTER_OPERATORS:
            raise DomainError("VALIDATION_ERROR", "Filtro inválido: revisa columna, operador y valor requeridos")
        if operator == "CONTAINS" and not isinstance(item.get("value"), str):
            raise DomainError("VALIDATION_ERROR", "CONTAINS requiere un valor de texto")
        if operator in {"GT", "GTE", "LT", "LTE"} and type_by_column[item["column"]] not in {
            "INTEGER", "NUMBER", "DATE", "DATETIME",
        }:
            raise DomainError("VALIDATION_ERROR", f"{operator} solo está disponible para columnas numéricas o de fecha")
        if operator in {"GT", "GTE", "LT", "LTE"}:
            kind, value = type_by_column[item["column"]], item.get("value")
            if kind in {"INTEGER", "NUMBER"} and _decimal(value) is None:
                raise DomainError("VALIDATION_ERROR", "La comparación numérica requiere un número finito")
            if kind == "DATE":
                try:
                    date.fromisoformat(value) if isinstance(value, str) else None
                except ValueError as exc:
                    raise DomainError("VALIDATION_ERROR", "El filtro DATE requiere una fecha ISO-8601") from exc
                if not isinstance(value, str):
                    raise DomainError("VALIDATION_ERROR", "El filtro DATE requiere una fecha ISO-8601")
            if kind == "DATETIME":
                try:
                    datetime.fromisoformat(value) if isinstance(value, str) else None
                except ValueError as exc:
                    raise DomainError("VALIDATION_ERROR", "El filtro DATETIME requiere una fecha ISO-8601") from exc
                if not isinstance(value, str):
                    raise DomainError("VALIDATION_ERROR", "El filtro DATETIME requiere una fecha ISO-8601")
        if "value" in item and isinstance(item["value"], (dict, list)):
            raise DomainError("VALIDATION_ERROR", "El valor del filtro debe ser escalar o null")
        if isinstance(item.get("value"), float) and not math.isfinite(item["value"]):
            raise DomainError("VALIDATION_ERROR", "El valor numérico del filtro debe ser finito")
        clean_filters.append({key: item[key] for key in sorted(item)})
    sort = spec.get("sort", [])
    if not isinstance(sort, list) or len(sort) > len(headers):
        raise DomainError("VALIDATION_ERROR", "sort debe ser una lista acotada de columnas")
    clean_sort = []
    seen_sort = set()
    for item in sort:
        if not isinstance(item, dict) or set(item) != {"column", "direction"}:
            raise DomainError("VALIDATION_ERROR", "Cada orden debe especificar column y direction")
        column, direction = item["column"], item["direction"]
        if not isinstance(column, str) or not isinstance(direction, str):
            raise DomainError("VALIDATION_ERROR", "Columna y dirección de sort deben ser texto")
        if column not in headers or direction not in _SORT_DIRECTIONS or column in seen_sort:
            raise DomainError("VALIDATION_ERROR", "Columna u orden de sort inválido o duplicado")
        seen_sort.add(column)
        clean_sort.append({"column": column, "direction": direction})
    return {"columns": columns[:], "filters": clean_filters, "sort": clean_sort, "limit": limit}


def _decimal(value):
    if value is None or value == "":
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return result if result.is_finite() else None


def _matches(row, rule, type_by_column):
    value = row[rule["column"]]
    operator = rule["operator"]
    empty = value is None or (isinstance(value, str) and not value.strip())
    if operator == "IS_EMPTY":
        return empty
    if operator == "NOT_EMPTY":
        return not empty
    expected = rule["value"]
    if operator == "CONTAINS":
        return not empty and str(expected).casefold() in str(value).casefold()
    if operator in {"EQ", "NE"}:
        kind = type_by_column[rule["column"]]
        if kind in {"INTEGER", "NUMBER"}:
            left, right = _decimal(value), _decimal(expected)
            equal = left is not None and right is not None and left == right
        else:
            equal = value == expected or (value is not None and expected is not None and str(value) == str(expected))
        return equal if operator == "EQ" else not equal
    if empty:
        return False
    kind = type_by_column[rule["column"]]
    if kind in {"INTEGER", "NUMBER"}:
        left, right = _decimal(value), _decimal(expected)
        if left is None or right is None:
            return False
    elif kind in {"DATE", "DATETIME"}:
        left, right = str(value), str(expected)
    else:
        raise DomainError("VALIDATION_ERROR", f"{operator} solo está disponible para columnas numéricas o de fecha")
    return {"GT": left > right, "GTE": left >= right, "LT": left < right, "LTE": left <= right}[operator]


def _sort_rows(rows, sort, type_by_column):
    def compare(left, right):
        for item in sort:
            a, b = left[item["column"]], right[item["column"]]
            a_empty = a is None or (isinstance(a, str) and not a.strip())
            b_empty = b is None or (isinstance(b, str) and not b.strip())
            if a_empty != b_empty:
                return 1 if a_empty else -1  # vacíos siempre al final
            if a_empty:
                continue
            if type_by_column[item["column"]] in {"INTEGER", "NUMBER"}:
                a, b = _decimal(a), _decimal(b)
                if a is None or b is None:
                    a, b = str(left[item["column"]]), str(right[item["column"]])
            else:
                a, b = str(a).casefold(), str(b).casefold()
            result = (a > b) - (a < b)
            if result:
                return -result if item["direction"] == "DESC" else result
        return 0

    return sorted(rows, key=cmp_to_key(compare)) if sort else rows


def _make_result(headers, source_rows, spec, version):
    type_by_column = {item["name"]: item["observed_type"] for item in version.columns_json}
    rows = [row for row in source_rows if all(_matches(row, rule, type_by_column) for rule in spec["filters"])]
    rows = _sort_rows(rows, spec["sort"], type_by_column)
    total_rows = len(rows)
    projected = [{column: row.get(column) for column in spec["columns"]} for row in rows[:spec["limit"]]]
    result = {
        "columns": spec["columns"],
        "rows": projected,
        "row_count": len(projected),
        "matched_row_count": total_rows,
        "truncated": total_rows > len(projected),
        "input": {"dataset_version_id": str(version.id), "version_no": version.version_no,
                  "sha256": version.sha256},
        "engine_version": ENGINE_VERSION,
    }
    encoded = _canonical_json(result)
    if len(encoded) > MAX_RESULT_BYTES:
        raise DomainError("OUTPUT_TOO_LARGE", "La tabla supera el máximo de 4 MiB; reduce columnas o limit")
    if not projected:
        raise DomainError("EMPTY_RESULT", "La receta no produjo filas; ajusta los filtros")
    return result, _sha256(encoded)


def _version_data(version):
    return {
        "id": str(version.id), "recipe_id": str(version.recipe_id), "version_no": version.version_no,
        "expected_input_sha256": version.expected_input_sha256, "spec": version.spec_json,
        "spec_sha256": version.spec_sha256, "result": version.result_json,
        "result_sha256": version.result_sha256, "engine_version": version.engine_version,
        "row_count": version.row_count, "column_count": version.column_count,
        "created_by_id": version.created_by_id, "created_at": version.created_at.isoformat(),
    }


def create_table_recipe(actor, project_id, dataset_id, dataset_version_no, name, spec,
                        idempotency_key, expected_version=None):
    project = require_project(actor, project_id, owner_only=True)
    try:
        key = UUID(str(idempotency_key))
    except (ValueError, TypeError, AttributeError) as exc:
        raise DomainError("VALIDATION_ERROR", "idempotency_key debe ser UUID") from exc
    if not isinstance(name, str) or not name.strip() or len(name.strip()) > 160:
        raise DomainError("VALIDATION_ERROR", "name debe contener entre 1 y 160 caracteres")
    input_version = ResearchDatasetVersion.objects.filter(
        dataset__project=project, dataset_id=dataset_id, version_no=dataset_version_no,
    ).select_related("dataset").first()
    if input_version is None:
        raise DomainError("NOT_FOUND", "Versión de dataset no encontrada")
    existing_identity = ResearchTableRecipe.objects.filter(dataset_version=input_version, name=name.strip()).first()
    if expected_version is not None and (type(expected_version) is not int or expected_version < 1):
        raise DomainError("VALIDATION_ERROR", "expected_version debe ser un entero positivo")
    if existing_identity is None and expected_version is not None:
        raise DomainError("VERSION_CONFLICT", "expected_version no aplica al crear una receta")
    if existing_identity is not None and expected_version is None:
        raise DomainError("VALIDATION_ERROR", "expected_version es obligatorio al versionar una receta existente")
    _, source_path = get_dataset_file(actor, project_id, dataset_id, dataset_version_no)
    headers, source_rows = _read_source_rows(source_path, input_version)
    type_by_column = {item["name"]: item["observed_type"] for item in input_version.columns_json}
    clean_spec = _validate_spec(spec, headers, type_by_column)
    spec_hash = _sha256(_canonical_json(clean_spec))
    result, result_hash = _make_result(headers, source_rows, clean_spec, input_version)
    try:
        with transaction.atomic():
            if existing_identity is None:
                recipe = ResearchTableRecipe.objects.create(
                    dataset_version=input_version, name=name.strip(), created_by=actor,
                )
            else:
                recipe = ResearchTableRecipe.objects.select_for_update().get(pk=existing_identity.pk)
            previous = recipe.versions.filter(idempotency_key=key).first()
            if previous:
                if previous.spec_sha256 != spec_hash or previous.expected_input_sha256 != input_version.sha256:
                    raise DomainError("VERSION_CONFLICT", "La clave idempotente ya se usó con otra receta")
                return {"recipe": _recipe_data(recipe), "version": _version_data(previous),
                        "created": False, "idempotent": True}
            latest = recipe.versions.order_by("-version_no").first()
            current_no = latest.version_no if latest else 0
            stale = current_no != expected_version if expected_version is not None else current_no != 0
            if stale:
                raise DomainError("VERSION_CONFLICT", "La receta cambió; vuelve a consultar la versión actual")
            version = ResearchTableRecipeVersion.objects.create(
                recipe=recipe, version_no=current_no + 1, idempotency_key=key,
                expected_input_sha256=input_version.sha256, spec_json=clean_spec, spec_sha256=spec_hash,
                result_json=result, result_sha256=result_hash, engine_version=ENGINE_VERSION,
                row_count=result["row_count"], column_count=len(result["columns"]), created_by=actor,
            )
    except IntegrityError as exc:
        raise DomainError("VERSION_CONFLICT", "La receta ya existe o tuvo un cambio concurrente") from exc
    return {"recipe": _recipe_data(recipe), "version": _version_data(version),
            "created": True, "idempotent": False}


def _recipe_data(recipe):
    latest = recipe.versions.order_by("-version_no").first()
    return {"id": str(recipe.id), "dataset_version_id": str(recipe.dataset_version_id),
            "name": recipe.name, "latest_version": latest.version_no if latest else None,
            "created_by_id": recipe.created_by_id, "created_at": recipe.created_at.isoformat()}


def list_table_recipes(actor, project_id, dataset_id, dataset_version_no):
    project = require_project(actor, project_id)
    input_version = ResearchDatasetVersion.objects.filter(
        dataset__project=project, dataset_id=dataset_id, version_no=dataset_version_no,
    ).first()
    if input_version is None:
        raise DomainError("NOT_FOUND", "Versión de dataset no encontrada")
    recipes = ResearchTableRecipe.objects.filter(dataset_version=input_version).prefetch_related("versions")
    return {"dataset_version_id": str(input_version.id), "recipes": [_recipe_data(item) for item in recipes]}


def list_table_recipe_versions(actor, project_id, dataset_id, recipe_id):
    project = require_project(actor, project_id)
    recipe = ResearchTableRecipe.objects.filter(
        pk=recipe_id, dataset_version__dataset__project=project, dataset_version__dataset_id=dataset_id,
    ).first()
    if recipe is None:
        raise DomainError("NOT_FOUND", "Receta no encontrada")
    return {"recipe": _recipe_data(recipe),
            "versions": [_version_data(item) for item in recipe.versions.order_by("-version_no")]}


def get_table_recipe_version(actor, project_id, dataset_id, recipe_id, version_no):
    project = require_project(actor, project_id)
    version = ResearchTableRecipeVersion.objects.filter(
        recipe_id=recipe_id, recipe__dataset_version__dataset__project=project,
        recipe__dataset_version__dataset_id=dataset_id, version_no=version_no,
    ).select_related("recipe").first()
    if version is None:
        raise DomainError("NOT_FOUND", "Versión de receta no encontrada")
    data = _version_data(version)
    if _sha256(_canonical_json(data["spec"])) != data["spec_sha256"]:
        raise DomainError("OUTPUT_INVALID", "El hash de la receta almacenada no coincide")
    if _sha256(_canonical_json(data["result"])) != data["result_sha256"]:
        raise DomainError("OUTPUT_INVALID", "El hash de la tabla almacenada no coincide")
    return {"recipe": _recipe_data(version.recipe), "version": data}
