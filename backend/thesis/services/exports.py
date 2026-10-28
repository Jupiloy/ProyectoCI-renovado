"""Exportación editable de las revisiones existentes del proyecto."""

from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt

from thesis.models import Chapter, ChapterRevision
from thesis.services.access import require_project
from thesis.services.errors import DomainError
from thesis.services.ieee_citations import build_ieee_citation_view
from thesis.services.steps import CHAPTER_STEPS
from thesis.step_catalogue import CATALOGO_PASOS


def _append_markdown(document: Document, content: str) -> None:
    """Conserva el texto redactado; solo interpreta encabezados y listas simples."""
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=4)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)


def _append_artifact(document: Document, block: dict, artifacts: dict) -> None:
    from docx.shared import Inches

    resolved = artifacts.get(block["id"])
    if resolved is None:
        raise DomainError("ARTIFACT_MANIFEST_INVALID", "El bloque no cuenta con un artefacto congelado")
    document.add_paragraph(block["caption"])
    if block["type"] == "data_table":
        result = resolved["result"]
        columns, rows = result["columns"], result["rows"]
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, columns):
            cell.text = str(label)
        for row in rows:
            cells = table.add_row().cells
            for cell, column in zip(cells, columns):
                value = row.get(column)
                cell.text = "" if value is None else str(value)
        if result.get("truncated"):
            document.add_paragraph(
                f"Nota: la receta coincidió con {result.get('matched_row_count')} filas; "
                f"se muestran {result.get('row_count')} filas congeladas."
            )
    else:
        paragraph = document.add_paragraph()
        run = paragraph.add_run()
        picture = run.add_picture(str(resolved["path"]), width=Inches(6.0))
        picture._inline.docPr.set("descr", resolved.get("accessibility", {}).get("alt_text", block["caption"]))


def _append_content_tree(document: Document, content_tree: dict, artifacts: dict | None = None) -> None:
    """Renderiza el árbol v1 congelado sin reinterpretar texto plano como sintaxis Markdown."""
    for block in content_tree.get("blocks", []):
        kind = block["type"]
        if kind == "heading":
            document.add_heading(block["text"], level=block["level"])
        elif kind == "paragraph":
            document.add_paragraph(block["text"])
        elif kind == "blockquote":
            document.add_paragraph(block["text"], style="Quote")
        elif kind == "code_block":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block["text"])
            run.font.name = "Courier New"
        elif kind == "bullet_list":
            for item in block["items"]:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "ordered_list":
            for item in block["items"]:
                document.add_paragraph(item, style="List Number")
        elif kind in {"data_table", "figure"}:
            _append_artifact(document, block, artifacts or {})


def _append_content_tree_with_ieee(document: Document, content_tree: dict, citation_view: dict,
                                   artifacts: dict | None = None) -> None:
    """Inserta marcadores IEEE en párrafos y añade referencias desde metadatos congelados."""
    marks_by_block = {}
    for mark in citation_view["marks"]:
        marks_by_block.setdefault(mark["block_id"], []).append(mark)
    for marks in marks_by_block.values():
        marks.sort(key=lambda mark: mark["position"])

    for block in content_tree.get("blocks", []):
        kind = block["type"]
        if kind == "paragraph" and marks_by_block.get(block["id"]):
            paragraph = document.add_paragraph()
            text = block["text"]
            cursor = 0
            for mark in marks_by_block[block["id"]]:
                position = mark["position"]
                paragraph.add_run(text[cursor:position])
                paragraph.add_run(mark["citation_text"])
                cursor = position
            paragraph.add_run(text[cursor:])
        elif kind == "paragraph":
            document.add_paragraph(block["text"])
        elif kind == "heading":
            document.add_heading(block["text"], level=block["level"])
        elif kind == "blockquote":
            document.add_paragraph(block["text"], style="Quote")
        elif kind == "code_block":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block["text"])
            run.font.name = "Courier New"
        elif kind == "bullet_list":
            for item in block["items"]:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "ordered_list":
            for item in block["items"]:
                document.add_paragraph(item, style="List Number")
        elif kind in {"data_table", "figure"}:
            _append_artifact(document, block, artifacts or {})

    if citation_view["references"]:
        document.add_heading("Referencias", level=1)
        document.add_paragraph(citation_view["notice"])
        for reference in citation_view["references"]:
            paragraph = document.add_paragraph(reference["ieee_text"])
            if reference["missing_fields"]:
                paragraph.add_run(
                    f" [Metadatos incompletos: {', '.join(reference['missing_fields'])}]"
                )


def export_draft_docx(actor, project_id) -> BytesIO:
    """Exporta el último texto guardado de cada paso como borrador, sin inventar contenido."""
    project = require_project(actor, project_id)
    revisions = ChapterRevision.objects.filter(chapter__project=project).select_related("chapter").order_by(
        "chapter__number", "step_code", "-revision_no"
    )
    latest = {}
    for revision in revisions:
        latest.setdefault(revision.step_code, revision)

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(3)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 2
    document.add_heading(project.title or project.name, level=0)
    document.add_paragraph("Borrador de trabajo generado a partir de las revisiones guardadas.")
    document.add_paragraph(
        "Revise las fuentes, citas, contenido y formato exigido por su institución antes de presentar este documento."
    )
    for chapter_no, codes in CHAPTER_STEPS.items():
        present = [code for code in codes if code in latest]
        if not present:
            continue
        document.add_heading(f"Capítulo {chapter_no}", level=1)
        for code in present:
            document.add_heading(f"{code}. {CATALOGO_PASOS[code]['nombre']}", level=2)
            _append_markdown(document, latest[code].content_md)
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def export_approved_snapshot_docx(actor, project_id, chapter_no: int) -> BytesIO:
    """Genera un DOCX desde el snapshot exacto aprobado; nunca usa revisiones actuales."""
    project = require_project(actor, project_id)
    chapter = Chapter.objects.select_related("approved_snapshot__document_revision").filter(
        project=project, number=chapter_no,
    ).first()
    if chapter is None:
        raise DomainError("NOT_FOUND", "Capítulo no encontrado")
    snapshot = chapter.approved_snapshot
    if chapter.status != Chapter.Status.APPROVED or snapshot is None:
        raise DomainError(
            "VERSION_CONFLICT",
            "El capítulo no tiene una aprobación vinculada a un snapshot exportable",
        )
    if snapshot.chapter_id != chapter.pk:
        raise DomainError("VERSION_CONFLICT", "El snapshot aprobado no pertenece al capítulo")

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(3)
    for style_name in (
        "Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6",
        "List Bullet", "List Number", "Quote",
    ):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 2

    document.add_heading(snapshot.project_title or "Proyecto sin título registrado en el snapshot", level=0)
    document.add_heading(f"Capítulo {chapter.number}", level=1)
    document.add_paragraph("Exportación congelada del snapshot aprobado.")
    document.add_paragraph(f"Snapshot: {snapshot.id} (n.º {snapshot.snapshot_no})")
    document.add_paragraph(f"SHA-256 del payload: {snapshot.payload_sha256}")
    if snapshot.document_revision_id:
        document.add_paragraph(f"Revisión del documento: {snapshot.document_revision_id}")
    if chapter.approved_at:
        document.add_paragraph(f"Aprobado en: {chapter.approved_at.isoformat()}")
    document.add_paragraph(
        "El texto exportado corresponde al contenido congelado del snapshot. "
        "Los outputs estructurados del manifest no se convierten automáticamente en prosa."
    )
    artifacts = {}
    if snapshot.document_revision_id and snapshot.document_artifacts_manifest:
        from thesis.services.document_artifacts import verify_snapshot_artifacts
        artifacts = verify_snapshot_artifacts(snapshot)
    if snapshot.document_revision_id:
        citation_view = build_ieee_citation_view(snapshot.document_revision)
        if citation_view["marks"]:
            _append_content_tree_with_ieee(document, snapshot.document_revision.content_tree, citation_view, artifacts)
        else:
            _append_content_tree(document, snapshot.document_revision.content_tree, artifacts)
    else:
        _append_markdown(document, snapshot.content_md)

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream
