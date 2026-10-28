"""Ingesta versionada de datasets CSV/XLSX sin ejecutar código del archivo."""

import csv
import hashlib
import math
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import UUID, uuid4

from django.conf import settings
from django.db import IntegrityError, transaction

from thesis.models import ResearchDataset, ResearchDatasetVersion
from thesis.services.access import require_project
from thesis.services.errors import DomainError


MAX_DATASET_BYTES = 30 * 1024 * 1024
MAX_DATASET_ROWS = 100_000
MAX_DATASET_COLUMNS = 256
MAX_XLSX_MEMBERS = 10_000
MAX_XLSX_EXPANDED_BYTES = 250 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 10_000
_HEADER_INVALID = re.compile(r"[\x00-\x1f]")


def _dataset_version_data(version: ResearchDatasetVersion) -> dict:
    return {
        "id": str(version.id), "dataset_id": str(version.dataset_id), "version_no": version.version_no,
        "original_filename": version.original_filename, "file_format": version.file_format,
        "sha256": version.sha256, "size_bytes": version.size_bytes,
        "row_count": version.row_count, "column_count": version.column_count,
        "columns": version.columns_json, "import_metadata": version.import_metadata_json,
        "created_by_id": version.created_by_id, "created_at": version.created_at.isoformat(),
    }


def _classify(value) -> str:
    if value is None or value == "":
        return "EMPTY"
    if isinstance(value, bool):
        return "BOOLEAN"
    if isinstance(value, (datetime, date)):
        return "DATETIME" if isinstance(value, datetime) else "DATE"
    if isinstance(value, int):
        return "INTEGER"
    if isinstance(value, (float, Decimal)):
        return "NUMBER" if math.isfinite(float(value)) else "TEXT"
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return "EMPTY"
        if re.fullmatch(r"[+-]?\d+", text):
            return "INTEGER"
        try:
            value = Decimal(text)
            return "NUMBER" if value.is_finite() else "TEXT"
        except InvalidOperation:
            pass
        try:
            date.fromisoformat(text)
            return "DATE"
        except ValueError:
            return "TEXT"
    return "TEXT"


def _finish_columns(headers, type_sets, blank_counts) -> list[dict]:
    columns = []
    for index, header in enumerate(headers):
        observed = sorted(type_sets[index] - {"EMPTY"})
        columns.append({
            "index": index,
            "name": header,
            "observed_type": observed[0] if len(observed) == 1 else ("MIXED" if observed else "UNKNOWN"),
            "nonempty_types": observed,
            "blank_cells": blank_counts[index],
        })
    return columns


def _validate_headers(raw_headers) -> list[str]:
    if not raw_headers:
        raise DomainError("INVALID_DATASET", "El archivo no contiene encabezados")
    if len(raw_headers) > MAX_DATASET_COLUMNS:
        raise DomainError("INVALID_DATASET", f"El dataset supera el límite de {MAX_DATASET_COLUMNS} columnas")
    headers = []
    for raw in raw_headers:
        name = str(raw).strip() if raw is not None else ""
        if not name or len(name) > 200 or _HEADER_INVALID.search(name):
            raise DomainError("INVALID_DATASET", "Cada columna requiere un encabezado legible de hasta 200 caracteres")
        headers.append(name)
    normalized = [name.casefold() for name in headers]
    if len(normalized) != len(set(normalized)):
        raise DomainError("INVALID_DATASET", "Los encabezados deben ser únicos, sin distinguir mayúsculas")
    return headers


def _inspect_csv(path: Path) -> dict:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream, strict=True)
            headers = _validate_headers(next(reader, []))
            types = [set() for _ in headers]
            blanks = [0 for _ in headers]
            rows = 0
            for row in reader:
                if len(row) > len(headers):
                    raise DomainError("INVALID_DATASET", "Una fila CSV tiene más valores que columnas")
                if not any(str(value).strip() for value in row):
                    continue
                rows += 1
                if rows > MAX_DATASET_ROWS:
                    raise DomainError("DATASET_TOO_LARGE", f"El dataset supera {MAX_DATASET_ROWS} filas")
                row.extend([""] * (len(headers) - len(row)))
                for index, value in enumerate(row):
                    kind = _classify(value)
                    types[index].add(kind)
                    if kind == "EMPTY":
                        blanks[index] += 1
    except UnicodeDecodeError as exc:
        raise DomainError("INVALID_DATASET", "El CSV debe estar codificado en UTF-8") from exc
    except csv.Error as exc:
        raise DomainError("INVALID_DATASET", "El CSV no tiene una estructura válida") from exc
    if rows == 0:
        raise DomainError("INVALID_DATASET", "El dataset debe contener al menos una fila de datos")
    return {
        "row_count": rows,
        "column_count": len(headers),
        "columns": _finish_columns(headers, types, blanks),
        "metadata": {"parser": "python.csv", "parser_version": "stdlib", "header_row": 1,
                     "blank_rows_ignored": True, "formula_cells": 0},
    }


def _validate_xlsx_package(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise DomainError("INVALID_DATASET", "El XLSX contiene demasiados componentes internos")
            expanded = sum(item.file_size for item in members)
            if expanded > MAX_XLSX_EXPANDED_BYTES:
                raise DomainError("INVALID_DATASET", "El contenido expandido del XLSX supera el límite local")
            for item in members:
                if item.file_size and item.file_size / max(item.compress_size, 1) > MAX_XLSX_COMPRESSION_RATIO:
                    raise DomainError("INVALID_DATASET", "El XLSX tiene una razón de compresión no admitida")
            names = {item.filename for item in members}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise DomainError("INVALID_DATASET", "El archivo no es un libro XLSX válido")
    except zipfile.BadZipFile as exc:
        raise DomainError("INVALID_DATASET", "El archivo XLSX está dañado") from exc


def _inspect_xlsx(path: Path, sheet_name: str | None) -> dict:
    _validate_xlsx_package(path)
    from openpyxl import load_workbook

    try:
        formulas_book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        values_book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise DomainError("INVALID_DATASET", "No se pudo leer la estructura XLSX") from exc
    try:
        visible = [sheet.title for sheet in formulas_book.worksheets if sheet.sheet_state == "visible"]
        if not visible:
            raise DomainError("INVALID_DATASET", "El libro no contiene hojas visibles")
        if sheet_name:
            if sheet_name not in visible:
                raise DomainError("VALIDATION_ERROR", "La hoja solicitada no existe o no es visible")
            selected = sheet_name
        elif len(visible) == 1:
            selected = visible[0]
        else:
            raise DomainError("VALIDATION_ERROR", "Indica sheet_name cuando el libro contiene varias hojas visibles")

        formula_sheet = formulas_book[selected]
        values_sheet = values_book[selected]
        if formula_sheet.max_column and formula_sheet.max_column > MAX_DATASET_COLUMNS:
            raise DomainError("DATASET_TOO_LARGE", f"La hoja supera el límite de {MAX_DATASET_COLUMNS} columnas")
        if formula_sheet.max_row and formula_sheet.max_row > MAX_DATASET_ROWS + 1:
            raise DomainError("DATASET_TOO_LARGE", f"La hoja supera el límite de {MAX_DATASET_ROWS} filas de datos")
        rows_iter = zip(formula_sheet.iter_rows(), values_sheet.iter_rows())
        first = next(rows_iter, None)
        if first is None:
            raise DomainError("INVALID_DATASET", "La hoja seleccionada está vacía")
        formula_header_cells, value_header_cells = first
        raw_headers = [cell.value for cell in value_header_cells]
        last_nonempty = 0
        for index, val in enumerate(raw_headers, start=1):
            if val not in (None, ""):
                last_nonempty = index
        if not last_nonempty:
            raise DomainError("INVALID_DATASET", "La primera fila de la hoja requiere encabezados")
        headers = _validate_headers(raw_headers[:last_nonempty])
        types = [set() for _ in headers]
        blanks = [0 for _ in headers]
        rows = 0
        formula_count = 0
        formula_without_cache = 0
        for formula_row, value_row in rows_iter:
            formula_row = formula_row[:last_nonempty]
            value_row = value_row[:last_nonempty]
            if not any(cell.value not in (None, "") for cell in value_row) and not any(
                cell.data_type == "f" for cell in formula_row
            ):
                continue
            rows += 1
            if rows > MAX_DATASET_ROWS:
                raise DomainError("DATASET_TOO_LARGE", f"El dataset supera {MAX_DATASET_ROWS} filas")
            for index, formula_cell in enumerate(formula_row):
                value_cell = value_row[index]
                if formula_cell.data_type == "f":
                    formula_count += 1
                    if value_cell.value is None:
                        formula_without_cache += 1
                kind = _classify(value_cell.value)
                types[index].add(kind)
                if kind == "EMPTY":
                    blanks[index] += 1
        if rows == 0:
            raise DomainError("INVALID_DATASET", "La hoja debe contener al menos una fila de datos")
        return {
            "row_count": rows,
            "column_count": len(headers),
            "columns": _finish_columns(headers, types, blanks),
            "metadata": {
                "parser": "openpyxl", "parser_version": __import__("openpyxl").__version__,
                "selected_sheet": selected, "visible_sheets": visible, "header_row": 1,
                "formula_cells": formula_count, "formula_cells_without_cached_value": formula_without_cache,
                "formula_values_recalculated": False,
                "analysis_warning": "Los valores de fórmula son cachés del libro y pueden estar desactualizados."
                    if formula_count else "",
            },
        }
    finally:
        formulas_book.close()
        values_book.close()


def _inspect(path: Path, extension: str, sheet_name: str | None) -> dict:
    if extension == ".csv":
        if sheet_name:
            raise DomainError("VALIDATION_ERROR", "sheet_name solo corresponde a XLSX")
        return {**_inspect_csv(path), "file_format": ResearchDatasetVersion.Format.CSV}
    if extension == ".xlsx":
        return {**_inspect_xlsx(path, sheet_name), "file_format": ResearchDatasetVersion.Format.XLSX}
    raise DomainError("UNSUPPORTED_FORMAT", "Los datasets admiten CSV UTF-8 y XLSX")


def _write_upload(uploaded_file, folder: Path) -> tuple[Path, str, int, str]:
    filename = Path(getattr(uploaded_file, "name", "")).name
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise DomainError("UNSUPPORTED_FORMAT", "Los datasets admiten CSV UTF-8 y XLSX")
    folder.mkdir(parents=True, exist_ok=True)
    # openpyxl valida la extensión del path; conservar la extensión validada en el staging.
    temporary = folder / f"upload-{uuid4().hex}{extension}"
    digest = hashlib.sha256()
    size = 0
    try:
        with temporary.open("wb") as target:
            for chunk in uploaded_file.chunks():
                size += len(chunk)
                if size > MAX_DATASET_BYTES:
                    raise DomainError("DATASET_TOO_LARGE", "El archivo supera el límite local de 30 MiB")
                target.write(chunk)
                digest.update(chunk)
        if size == 0:
            raise DomainError("INVALID_DATASET", "El archivo está vacío")
        if extension == ".xlsx":
            with temporary.open("rb") as raw:
                if raw.read(2) != b"PK":
                    raise DomainError("INVALID_DATASET", "El archivo XLSX no tiene cabecera ZIP")
        return temporary, digest.hexdigest(), size, extension
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _parse_key(value) -> UUID:
    try:
        return UUID(str(value))
    except (ValueError, TypeError, AttributeError) as exc:
        raise DomainError("VALIDATION_ERROR", "idempotency_key debe ser UUID") from exc


def _present_dataset(dataset: ResearchDataset) -> dict:
    latest = dataset.versions.order_by("-version_no").first()
    return {
        "id": str(dataset.id), "project_id": str(dataset.project_id), "name": dataset.name,
        "description": dataset.description, "created_by_id": dataset.created_by_id,
        "created_at": dataset.created_at.isoformat(),
        "latest_version": _dataset_version_data(latest) if latest else None,
    }


def _create_version(dataset, actor, temp, digest, size, filename, extension, info, key, expected_version=None):
    with transaction.atomic():
        locked = ResearchDataset.objects.select_for_update().get(pk=dataset.pk)
        existing = locked.versions.filter(idempotency_key=key).first()
        if existing:
            if existing.sha256 != digest:
                raise DomainError("VERSION_CONFLICT", "La clave idempotente ya se usó con otro archivo")
            return existing, False
        current = locked.versions.order_by("-version_no").first()
        current_no = current.version_no if current else 0
        if expected_version is not None and expected_version != current_no:
            raise DomainError("VERSION_CONFLICT", "El dataset cambió; vuelve a consultar su versión actual")
        next_no = current_no + 1
        root = Path(settings.MEDIA_ROOT).resolve()
        destination = root / "datasets" / str(locked.project_id) / str(locked.id) / f"v{next_no}-{digest[:12]}{extension}"
        destination.parent.mkdir(parents=True, exist_ok=True)
        version = ResearchDatasetVersion.objects.create(
            dataset=locked, version_no=next_no, idempotency_key=key, original_filename=filename,
            file_format=info["file_format"], storage_key=str(destination.relative_to(root)), sha256=digest,
            size_bytes=size, row_count=info["row_count"], column_count=info["column_count"],
            columns_json=info["columns"], import_metadata_json=info["metadata"], created_by=actor,
        )
        temp.replace(destination)
        return version, True


def create_dataset(actor, project_id, name: str, uploaded_file, idempotency_key, *, description="", sheet_name=None):
    project = require_project(actor, project_id, owner_only=True)
    name = str(name or "").strip()
    description = str(description or "").strip()
    if not name or len(name) > 160 or len(description) > 4000:
        raise DomainError("VALIDATION_ERROR", "Nombre (1–160) y descripción (hasta 4000 caracteres) requeridos")
    key = _parse_key(idempotency_key)
    folder = Path(settings.MEDIA_ROOT) / "datasets" / str(project.id) / "staging"
    temp, digest, size, extension = _write_upload(uploaded_file, folder)
    try:
        info = _inspect(temp, extension, sheet_name)
        existing_version = ResearchDatasetVersion.objects.filter(dataset__project=project, idempotency_key=key).first()
        if existing_version:
            if existing_version.sha256 != digest:
                raise DomainError("VERSION_CONFLICT", "La clave idempotente ya se usó con otro archivo")
            dataset = existing_version.dataset
            return {"dataset": _present_dataset(dataset), "version": _dataset_version_data(existing_version),
                    "created": False, "idempotent": True}
        try:
            with transaction.atomic():
                dataset = ResearchDataset.objects.create(project=project, name=name, description=description,
                                                         created_by=actor)
                version, created = _create_version(dataset, actor, temp, digest, size,
                    Path(getattr(uploaded_file, "name", "dataset" )).name, extension, info, key)
        except IntegrityError as exc:
            raise DomainError("VERSION_CONFLICT", "Ya existe un dataset con ese nombre en el proyecto") from exc
        return {"dataset": _present_dataset(dataset), "version": _dataset_version_data(version),
                "created": created, "idempotent": not created}
    finally:
        temp.unlink(missing_ok=True)


def add_dataset_version(actor, project_id, dataset_id, uploaded_file, idempotency_key, expected_version: int,
                        *, sheet_name=None):
    project = require_project(actor, project_id, owner_only=True)
    dataset = ResearchDataset.objects.filter(project=project, pk=dataset_id).first()
    if dataset is None:
        raise DomainError("NOT_FOUND", "Dataset no encontrado")
    if type(expected_version) is not int or expected_version < 1:
        raise DomainError("VALIDATION_ERROR", "expected_version debe ser la versión actual positiva")
    key = _parse_key(idempotency_key)
    folder = Path(settings.MEDIA_ROOT) / "datasets" / str(project.id) / str(dataset.id) / "staging"
    temp, digest, size, extension = _write_upload(uploaded_file, folder)
    try:
        info = _inspect(temp, extension, sheet_name)
        version, created = _create_version(dataset, actor, temp, digest, size,
            Path(getattr(uploaded_file, "name", "dataset")).name, extension, info, key, expected_version)
        dataset.refresh_from_db()
        return {"dataset": _present_dataset(dataset), "version": _dataset_version_data(version),
                "created": created, "idempotent": not created}
    finally:
        temp.unlink(missing_ok=True)


def list_datasets(actor, project_id):
    project = require_project(actor, project_id)
    datasets = ResearchDataset.objects.filter(project=project).prefetch_related("versions").order_by("created_at", "id")
    return {"datasets": [_present_dataset(dataset) for dataset in datasets]}


def list_dataset_versions(actor, project_id, dataset_id):
    project = require_project(actor, project_id)
    dataset = ResearchDataset.objects.filter(project=project, pk=dataset_id).first()
    if dataset is None:
        raise DomainError("NOT_FOUND", "Dataset no encontrado")
    versions = dataset.versions.select_related("created_by").order_by("-version_no")
    return {"dataset": {"id": str(dataset.id), "name": dataset.name},
            "versions": [_dataset_version_data(version) for version in versions]}


def get_dataset_version(actor, project_id, dataset_id, version_no: int, *, include_path=False):
    project = require_project(actor, project_id)
    version = ResearchDatasetVersion.objects.filter(
        dataset__project=project, dataset_id=dataset_id, version_no=version_no,
    ).select_related("dataset").first()
    if version is None:
        raise DomainError("NOT_FOUND", "Versión de dataset no encontrada")
    result = {"dataset": {"id": str(version.dataset_id), "name": version.dataset.name},
              "version": _dataset_version_data(version)}
    if include_path:
        root = Path(settings.MEDIA_ROOT).resolve()
        path = (root / version.storage_key).resolve()
        if root not in path.parents or not path.is_file():
            raise DomainError("NOT_FOUND", "Archivo de dataset no disponible")
        result["path"] = path
    return result


def get_dataset_file(actor, project_id, dataset_id, version_no: int) -> tuple[ResearchDatasetVersion, Path]:
    result = get_dataset_version(actor, project_id, dataset_id, version_no, include_path=True)
    version = ResearchDatasetVersion.objects.get(pk=result["version"]["id"])
    path = result["path"]
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != version.sha256:
        raise DomainError("OUTPUT_INVALID", "El hash del archivo almacenado no coincide con su versión")
    return version, path
