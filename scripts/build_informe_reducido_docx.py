#!/usr/bin/env python3
"""
Compilador del Informe Base de la Versión Reducida (07_INFORME_BASE_SEMANA_10.md)
a documento Microsoft Word (.docx) con Maquetación IEEE Profesional.

Cumple con:
- Doble columna balanceada (two-column IEEE layout).
- Tipografía estricta Times New Roman.
- Portada y Resumen Ejecutivo / Abstract en una sola columna con metadatos institucionales.
- Tablas y Figuras a ancho completo (17.2 cm) con saltos de sección continuos.
- Renderizado de todos los diagramas oficiales (00 a 07) y flujos operativos.
- Citas numéricas y referencias bibliográficas estructuradas.
- Paginación dinámica en pie de página.
"""

from __future__ import annotations

import os
import re
import subprocess
from pathlib import Path
from PIL import Image

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC_MD = ROOT / "V1" / "docs" / "07_INFORME_BASE_SEMANA_10.md"
DOCX_OUT = ROOT / "docs" / "06_informe_final_curso" / "informe_version_reducida_IEEE.docx"
DIAGRAMS_DIR = ROOT / "V1" / "docs" / "diagramas"
TEMP_DIR = Path("/tmp/tesis_mcp_docx_ieee")
TEMP_DIR.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D1D5DB", sz="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        borders.append(el)


def set_cell_margins(cell, top=70, bottom=70, left=90, right=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)


def set_section(section, columns: int = 2, top_cm=1.7, bottom_cm=1.7, left_cm=1.5, right_cm=1.5):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(top_cm)
    section.bottom_margin = Cm(bottom_cm)
    section.left_margin = Cm(left_cm)
    section.right_margin = Cm(right_cm)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)
    cols = section._sectPr.xpath("./w:cols")
    cols = cols[0] if cols else OxmlElement("w:cols")
    cols.set(qn("w:num"), str(columns))
    cols.set(qn("w:space"), "340")
    if cols.getparent() is None:
        section._sectPr.append(cols)


def add_formatted_runs(p, source: str) -> None:
    source = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", source)
    source = source.replace("<br>", "\n").replace("<br/>", "\n")
    source = source.replace("\\rightarrow", "→").replace("\\ge", "≥").replace("\\le", "≤")
    
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*)|\$[^$]+\$)", source)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = part.startswith("*") and part.endswith("*") and not bold
        code = part.startswith("`") and part.endswith("`")
        math = part.startswith("$") and part.endswith("$")
        
        if bold:
            val = part[2:-2]
        elif italic or code or math:
            val = part[1:-1]
        else:
            val = part
            
        run = p.add_run(val)
        run.bold = bold
        run.italic = italic or math
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)


def set_font(style, size: float, bold: bool = False, italic: bool = False) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)


def render_sequence_diagram() -> Path:
    seq_png = DIAGRAMS_DIR / "08_secuencia_guardar_paso.png"
    if seq_png.exists():
        return seq_png
    out_png = TEMP_DIR / "sequence_6_4_1.png"
    dot_content = """digraph Sequence_Paso {
        graph [bgcolor="white", fontname="Arial", nodesep=0.4, ranksep=0.45, pad="0.2"];
        node [shape=box, style="rounded", fillcolor="#ffffff", color="#000000", fontname="Arial", fontsize=9.5, margin="0.12,0.08"];
        edge [color="#000000", arrowsize=0.6, fontname="Arial", fontsize=8.5];
        p1 [label="1. Tesista confirma salida del paso en Cliente MCP / API REST"];
        p2 [label="2. Cliente envía payload, contrato Pydantic y lock_version esperada"];
        p3 [label="3. StepService valida autenticación, permisos y SalidaEstricta (extra='forbid')"];
        p4 [label="4. StepService bloquea proyecto en PostgreSQL y verifica lock_version concurrente"];
        p5 [label="5. Transacción atómica: materializa entidades, crea revisión e incrementa lock_version"];
        p6 [label="6. Devuelve confirmación, revision_id y nueva lock_version al cliente"];
        fail [label="Si lock_version difiere: VERSION_CONFLICT | Si contrato falla: VALIDATION_ERROR (Rollback)"];
        p1 -> p2 -> p3 -> p4 -> p5 -> p6;
        p4 -> fail [style=dashed, label="Conflicto"];
        p3 -> fail [style=dashed, label="Error contrato"];
    }"""
    dot_file = TEMP_DIR / "sequence.dot"
    dot_file.write_text(dot_content, encoding="utf-8")
    subprocess.run(["dot", "-Tpng", "-Gdpi=180", str(dot_file), "-o", str(out_png)], check=True)
    return out_png


def render_state_diagram() -> Path:
    state_png = DIAGRAMS_DIR / "09_estados_revision_capitulo.png"
    if state_png.exists():
        return state_png
    out_png = TEMP_DIR / "state_6_4_2.png"
    dot_content = """digraph State_Revision {
        graph [bgcolor="white", fontname="Arial", nodesep=0.5, ranksep=0.55, pad="0.2"];
        node [shape=box, style="rounded", fillcolor="#ffffff", color="#000000", fontname="Arial", fontsize=9.5, margin="0.14,0.08"];
        edge [color="#000000", arrowsize=0.65, fontname="Arial", fontsize=8.5];
        draft [label="DRAFT\\n(Borrador en desarrollo por Tesista)"];
        ready [label="READY\\n(Capítulo completado y enviado a revisión)"];
        changes [label="CHANGES_REQUESTED\\n(Asesor registra observaciones ancladas)"];
        approved [label="APPROVED\\n(Aprobado oficialmente sin observaciones abiertas)"];
        draft -> ready [label="Tesista envía capítulo\\n[Guarda: P01-P12 completados]"];
        ready -> changes [label="Asesor solicita cambios\\n[Registra observaciones]"];
        changes -> draft [label="Tesista subsana\\n[Crea nueva revisión inmutable]"];
        ready -> approved [label="Asesor aprueba capítulo\\n[Guarda: 0 observaciones abiertas]"];
        approved -> draft [label="Asesor reabre capítulo\\n[Solo Asesor y con motivo]", style=dashed];
    }"""
    dot_file = TEMP_DIR / "state.dot"
    dot_file.write_text(dot_content, encoding="utf-8")
    subprocess.run(["dot", "-Tpng", "-Gdpi=180", str(dot_file), "-o", str(out_png)], check=True)
    return out_png


def insert_full_width_image(doc: Document, img_path: Path, caption_text: str) -> None:
    # 1. Switch to 1 column for full width display
    set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), columns=1)
    
    # Optimize image for Word embedding (high crispness, < 100 KB each)
    opt_dir = TEMP_DIR / "opt"
    opt_dir.mkdir(exist_ok=True)
    opt_path = opt_dir / img_path.name
    
    with Image.open(img_path) as im:
        w_px, h_px = im.size
        # Resize if wider than 2000px for print publication crispness
        if w_px > 2000:
            target_h = int(h_px * (2000 / w_px))
            im_resized = im.resize((2000, target_h), Image.Resampling.LANCZOS)
        else:
            im_resized = im
            
        im_opt = im_resized.convert("RGB").convert("P", palette=Image.ADAPTIVE, colors=32)
        im_opt.save(str(opt_path), optimize=True)
        w_px, h_px = im_resized.size
    
    max_w_cm = 17.2
    max_h_cm = 16.0
    display_w = min(max_w_cm, max_h_cm * w_px / h_px)
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    p_img.add_run().add_picture(str(opt_path), width=Cm(display_w))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    run_cap = p_cap.add_run(caption_text)
    run_cap.italic = True
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(8.5)
    
    # 2. Switch back to 2 columns
    set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), columns=2)


def insert_full_width_table(doc: Document, rows_raw: list[str]) -> None:
    # 1. Switch to 1 column for full width table
    set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), columns=1)
    
    rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows_raw]
    rows = [rows[0]] + rows[2:]  # omit separator row
    
    cols_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols_count)
    table.alignment = 1  # center
    table.autofit = False
    
    # Calculate widths based on content length
    col_lens = [max(len(row[c]) if c < len(row) else 1 for row in rows) for c in range(cols_count)]
    total_len = sum(col_lens)
    TOTAL_WIDTH_CM = 17.2
    col_widths = [Cm(TOTAL_WIDTH_CM * (l / total_len)) for l in col_lens]
    
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols_count):
            cell = table.cell(r_idx, c_idx)
            cell.width = col_widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell, color="D1D5DB", sz="4")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            
            # Shading: Header vs Data
            if r_idx == 0:
                set_cell_shading(cell, "E9EDF1")
            elif r_idx % 2 == 1:
                set_cell_shading(cell, "FFFFFF")
            else:
                set_cell_shading(cell, "F8F9FA")
                
            text = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.02
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            
            add_formatted_runs(p, text)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(8.5)
                if r_idx == 0:
                    r.bold = True
                    
    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
    
    # 2. Switch back to 2 columns
    set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), columns=2)


def build_ieee_docx(input_path: Path = SRC_MD, output_path: Path = DOCX_OUT) -> Path:
    doc = Document()
    set_section(doc.sections[0], columns=1)  # Front matter in 1 column
    
    styles = doc.styles
    set_font(styles["Normal"], 10)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Normal"].paragraph_format.space_after = Pt(3)
    styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    for s_name, size in [("Heading 1", 11), ("Heading 2", 10.5), ("Heading 3", 10)]:
        set_font(styles[s_name], size, bold=True)
        styles[s_name].paragraph_format.space_before = Pt(8)
        styles[s_name].paragraph_format.space_after = Pt(3)
        styles[s_name].paragraph_format.keep_with_next = True
        
    # Page numbering in footer
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    
    lines = input_path.read_text(encoding="utf-8").splitlines()
    
    in_two_columns = False
    in_frontmatter = True
    in_abstract = False
    in_code_block = False
    code_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Horizontal dividers
        if stripped == "---":
            i += 1
            continue
            
        # Empty lines
        if not stripped:
            i += 1
            continue
            
        # Code blocks
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                lang = line[3:].strip()
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                # Process code block
                # Check if it was a Mermaid block
                mermaid_text = "\n".join(code_lines)
                if "sequenceDiagram" in mermaid_text:
                    seq_png = render_sequence_diagram()
                    insert_full_width_image(doc, seq_png, "Figura 6.3: Diagrama de Secuencia del Guardado de un Paso y Concurrencia Optimista.")
                elif "stateDiagram-v2" in mermaid_text:
                    state_png = render_state_diagram()
                    insert_full_width_image(doc, state_png, "Figura 6.4: Diagrama de Estados de la Máquina de Revisión Docente (Capitulo.status).")
                elif "flowchart TD" in mermaid_text:
                    arch_png = DIAGRAMS_DIR / "00_arquitectura_sistema.png"
                    insert_full_width_image(doc, arch_png, "Figura 6.5: Arquitectura del Sistema TesisMCP (Capas de Cliente, Transporte, Dominio y Persistencia).")
                else:
                    # Generic code box
                    p_box = doc.add_paragraph()
                    p_box.paragraph_format.left_indent = Cm(0.3)
                    p_box.paragraph_format.right_indent = Cm(0.3)
                    p_box.paragraph_format.space_before = Pt(3)
                    p_box.paragraph_format.space_after = Pt(3)
                    p_box.paragraph_format.line_spacing = 1.0
                    for c_line in code_lines:
                        run = p_box.add_run(c_line + "\n")
                        run.font.name = "Consolas"
                        run.font.size = Pt(8.5)
                code_lines = []
                i += 1
                continue
                
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Blockquotes (callouts)
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_runs(p, line[2:].strip())
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(9)
            i += 1
            continue
            
        # Images: ![alt](path)
        img_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if img_match:
            alt, rel_path = img_match.groups()
            img_file = (SRC_MD.parent / rel_path).resolve()
            # Check caption on subsequent non-empty line
            caption = f"Figura: {alt}"
            next_idx = i + 1
            while next_idx < len(lines) and not lines[next_idx].strip():
                next_idx += 1
            if next_idx < len(lines) and (lines[next_idx].strip().startswith("*Figura") or lines[next_idx].strip().startswith("Figura")):
                caption = lines[next_idx].strip().strip("*").strip()
                i = next_idx
            if img_file.exists():
                insert_full_width_image(doc, img_file, caption)
            else:
                p = doc.add_paragraph(f"[Imagen no encontrada: {rel_path}]")
            i += 1
            continue
            
        # Markdown Tables: | col1 | col2 |
        if stripped.startswith("|"):
            tbl_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_rows.append(lines[i].strip())
                i += 1
            insert_full_width_table(doc, tbl_rows)
            continue
            
        # Headings
        if line.startswith("# "):
            # Main Document Title
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("DESARROLLO DE UNA PLATAFORMA COLABORATIVA BASADA EN EL PROTOCOLO MCP, ROLES SECUENCIALES Y VALIDACIÓN DETERMINISTA PARA LA ASISTENCIA METODOLÓGICA Y SUPERVISIÓN DE TESIS UNIVERSITARIAS\n")
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(12)
            run_sub = p_sub.add_run("Informe Base del Proyecto — Curso Integrador I: Sistemas - Software (Capítulos 1 al 3)")
            run_sub.italic = True
            run_sub.font.name = "Times New Roman"
            run_sub.font.size = Pt(11)
            i += 1
            continue
            
        if line.startswith("## ESCUELA") or line.startswith("### TÍTULO"):
            # Omit redundant institution title header lines as they are merged into title block
            i += 1
            continue
            
        if line.startswith("- **Autor") or line.startswith("- **Integrantes") or line.startswith("- **Institución") or line.startswith("- **Curso") or line.startswith("- **Fecha"):
            # Institutional metadata
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_formatted_runs(p, line[2:].strip())
            for r in p.runs:
                r.font.size = Pt(9.5)
            i += 1
            continue
            
        if line.startswith("## RESUMEN EJECUTIVO"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("RESUMEN EJECUTIVO\n")
            run.bold = True
            run.font.size = Pt(10.5)
            in_abstract = True
            i += 1
            continue
            
        if in_abstract and line.startswith("**Palabras clave:"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(14)
            add_formatted_runs(p, line.strip())
            for r in p.runs:
                r.font.size = Pt(9.5)
            in_abstract = False
            
            # Switch to 2 columns after Abstract
            if not in_two_columns:
                set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), columns=2)
                in_two_columns = True
            i += 1
            continue
            
        if line.startswith("## "):
            title = line[3:].strip()
            
            # Roman numerals mapping for main chapters
            roman_map = {
                "QUÉ PUEDE HACER EL SISTEMA HOY": "I. QUÉ PUEDE HACER EL SISTEMA HOY",
                "ÍNDICE PROPUESTO": "II. ESTRUCTURA Y CATÁLOGO DE LA VERSIÓN REDUCIDA",
                "CAPÍTULO I: DESCRIPCIÓN DE LA SITUACIÓN PROBLEMÁTICA": "III. DESCRIPCIÓN DE LA SITUACIÓN PROBLEMÁTICA",
                "CAPÍTULO II: MARCO CONCEPTUAL Y ANTECEDENTES": "IV. MARCO CONCEPTUAL Y ANTECEDENTES",
                "CAPÍTULO III: OBJETIVOS Y JUSTIFICACIÓN": "V. OBJETIVOS Y JUSTIFICACIÓN",
                "CAPÍTULO IV: PLANIFICACIÓN DEL PROYECTO": "VI. PLANIFICACIÓN DEL PROYECTO",
                "CAPÍTULO V: TÉCNICAS E INSTRUMENTOS DE RECOLECCIÓN DE DATOS": "VII. TÉCNICAS E INSTRUMENTOS DE RECOLECCIÓN DE DATOS",
                "CAPÍTULO VI: METODOLOGÍA Y DESARROLLO DE LA SOLUCIÓN": "VIII. METODOLOGÍA Y DESARROLLO DE LA SOLUCIÓN",
                "CONCLUSIONES PRELIMINARES": "IX. CONCLUSIONES PRELIMINARES",
                "RECOMENDACIONES": "X. RECOMENDACIONES",
                "REFERENCIAS BIBLIOGRÁFICAS BASE": "REFERENCIAS BIBLIOGRÁFICAS",
                "ANEXO A: LISTA DE CAMPOS POR COMPLETAR ANTES DE PRESENTAR": "ANEXO A: CAMPOS PENDIENTES DE PRESENTACIÓN",
                "ANEXO B: GLOSARIO BREVE": "ANEXO B: GLOSARIO BREVE DE TÉRMINOS",
            }
            formatted_title = roman_map.get(title, title)
            
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.space_before = Pt(11)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_runs(p, formatted_title)
            i += 1
            continue
            
        if line.startswith("### "):
            title = line[4:].strip()
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, title)
            i += 1
            continue
            
        if line.startswith("#### "):
            title = line[5:].strip()
            p = doc.add_paragraph(style="Heading 3")
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, title)
            i += 1
            continue
            
        # Bullet list or numbered list
        if re.match(r"^(?:\*|\d+\.)\s", stripped):
            indent_level = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.45 + indent_level * 0.35)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            
            is_bullet = stripped.startswith("*")
            prefix = "• " if is_bullet else re.match(r"^\d+\.", stripped).group(0) + " "
            item_text = re.sub(r"^(?:\*|\d+\.)\s+", "", stripped)
            add_formatted_runs(p, prefix + item_text)
            i += 1
            continue
            
        # References with [N]
        if re.match(r"^\[\d+\] ", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.02
            add_formatted_runs(p, stripped)
            for r in p.runs:
                r.font.size = Pt(9)
            i += 1
            continue
            
        # Standard paragraph
        p = doc.add_paragraph(style="Normal")
        add_formatted_runs(p, line.strip())
        i += 1
        
    # Final section to balance columns
    set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), columns=1)
    
    doc.core_properties.title = "Informe del Proyecto TesisMCP (Edición Capítulos 1 al 3)"
    doc.core_properties.subject = "Curso Integrador I: Sistemas - Software"
    doc.core_properties.author = "Josué Carrasco García"
    
    doc.save(str(output_path))
    print(f"Documento DOCX IEEE compilado con éxito: {output_path} ({output_path.stat().st_size // 1024} KB)")
    return output_path


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Compilador del Informe Base de la Versión Reducida a DOCX IEEE.")
    parser.add_argument("--input", "-i", type=Path, default=SRC_MD, help="Ruta al archivo Markdown fuente.")
    parser.add_argument("--output", "-o", type=Path, default=DOCX_OUT, help="Ruta al archivo DOCX de salida.")
    args = parser.parse_args()

    out = build_ieee_docx(input_path=args.input, output_path=args.output)
