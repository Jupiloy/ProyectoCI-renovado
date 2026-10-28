"""Build a separate IEEE edition of the existing Curso Integrador report."""

from __future__ import annotations

import re
import subprocess
from html import escape
from pathlib import Path

from PIL import Image

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "06_informe_final_curso"
MD_OUT = OUT / "informe_academico_curso_integrador_IEEE.md"
DOCX_OUT = OUT / "informe_academico_curso_integrador_IEEE.docx"
DIAGRAM_SOURCE_OUT = ROOT / "docs" / "02_arquitectura_y_sistema" / "diagramas_editables"
TEMP = Path("/tmp/tesis_mcp_ieee")
TEMP.mkdir(exist_ok=True)


def ieee_markdown() -> str:
    """Read the maintained IEEE Markdown source for the Word edition."""
    return MD_OUT.read_text(encoding="utf-8")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)


def set_cell_margin(cell, value: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)


def clean_inline(value: str) -> str:
    value = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", value)
    value = value.replace("<br>", "\n").replace("$", "")
    value = value.replace("**", "").replace("`", "")
    value = re.sub(r"(?<!\w)\*([^*]+)\*", r"\1", value)
    return value.strip()


def add_formatted(p, source: str) -> None:
    source = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", source)
    source = source.replace("<br>", "\n").replace("$", "").replace("\\rightarrow", "→")
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))", source)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = part.startswith("*") and part.endswith("*") and not bold
        code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if bold else part[1:-1] if italic or code else part
        value = value.replace("**", "")
        run = p.add_run(value)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def set_font(style, size: float, bold: bool = False) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def diagram_png(which: int) -> Path:
    if which == 2:
        seq_p = ROOT / "V1" / "docs" / "diagramas" / "08_secuencia_guardar_paso.png"
        if seq_p.exists():
            return seq_p
    if which == 3:
        state_p = ROOT / "V1" / "docs" / "diagramas" / "09_estados_revision_capitulo.png"
        if state_p.exists():
            return state_p

    base = r'''
      graph [bgcolor="white", fontname="Arial", nodesep=0.40, ranksep=0.60];
      node [shape=box, style="rounded,filled", fillcolor="#F4F6F8", color="#8A9BAA", fontname="Arial", fontsize=10, margin="0.12,0.08"];
      edge [color="#526477", arrowsize=0.65, fontname="Arial", fontsize=9];
    '''
    diagrams = {
        1: r'''digraph G {''' + base + r'''
          rankdir=TB;
          start [label="Crear proyecto"]; step [label="Elegir paso P01–P32"];
          gate [shape=diamond,label="¿Prerequisitos?",fillcolor="#FFF5E8"];
          blocked [label="Mostrar faltantes\nSin cambios"];
          context [label="Preparar contexto\ny fuentes"];
          edit [label="Tesista revisa\ny edita propuesta"];
          valid [shape=diamond,label="¿Versión y reglas válidas?",fillcolor="#FFF5E8"];
          error [label="Conflicto o errores\nSin guardar"];
          rev [label="Crear revisión inmutable\ny actualizar estado"];
          submit [label="Enviar capítulo"]; advisor [label="Asesor revisa"];
          approve [shape=diamond,label="¿Aprueba?",fillcolor="#FFF5E8"];
          comments [label="Comentarios\ny solicitud de cambios"];
          done [label="Aprobar revisión\ny desbloquear siguiente"];
          start->step->gate; gate->blocked[label="No"]; gate->context[label="Sí"];
          context->edit->valid; valid->error[label="No"]; valid->rev[label="Sí"];
          rev->submit->advisor->approve; approve->comments[label="No"];
          approve->done[label="Sí"]; comments->edit;
        }''',
        2: r'''digraph G {''' + base + r'''
          rankdir=TB;
          p1 [label="1  Tesista → API/MCP\npreparar_paso(project_id, step_code)"];
          p2 [label="2  StepService → DB\nPermiso, prerequisitos y state_version"];
          p3 [label="3  Servicio → cliente\nPlantilla, fuentes y versión actual"];
          p4 [label="4  Tesista → API/MCP\nTexto confirmado + expected_version"];
          p5 [label="5  Servicio → DB\nValidación y transacción atómica"];
          p6 [label="6  Servicio → cliente\nrevision_id + versión nueva"];
          fail [label="Si falla: PASO_BLOQUEADO o VERSION_CONFLICT\nNo se modifica el proyecto",fillcolor="#FFF0F0"];
          p1->p2->p3->p4->p5->p6;
          p2->fail [style=dashed,label="Bloqueado"];
          p5->fail [style=dashed,label="Conflicto"];
        }''',
        3: r'''digraph G {''' + base + r'''
          rankdir=TB;
          draft [label="DRAFT\nBorrador"];
          ready [label="READY\nEnviado"];
          changes [label="CHANGES_REQUESTED\nCon observaciones"];
          approved [label="APPROVED\nAprobado"];
          draft->ready [label="Tesista envía revisión vigente"];
          ready->changes [label="Asesor observa"];
          changes->draft [label="Tesista crea nueva revisión"];
          ready->approved [label="Asesor aprueba sin pendientes"];
          approved->draft [label="Asesor reabre con motivo"];
        }''',
        4: r'''digraph G {''' + base + r'''
          rankdir=TB; splines=polyline;
          user [label="USER\nid PK; email UK"];
          member [label="PROJECT_PARTICIPANT\nproject_id + user_id UK\nrole; active"];
          project [label="PROJECT\nid UUID PK; title; status"];
          profile [label="RESEARCH_PROFILE\nproject_id PK/FK\ntopic; problem_statement; lock_version"];
          step [label="STEP_PROGRESS\nproject_id + step_code UK\nstatus; output_json"];
          chapter [label="CHAPTER\nproject_id + number UK\nstatus; approved_revision_id FK"];
          user->member; project->member;
          project->profile; project->step; project->chapter;
          {rank=same; profile; step; chapter;}
        }''',
        5: r'''digraph G {''' + base + r'''
          rankdir=TB; splines=polyline;
          user [label="USER\nid PK"];
          chapter [label="CHAPTER\nid PK; status\napproved_revision_id FK"];
          revision [label="CHAPTER_REVISION\nchapter_id + revision_no UK\nidempotency_key UK; content_md"];
          comment [label="REVIEW_COMMENT\nrevision_id FK; author_id FK\nstart_cp; end_cp; status"];
          event [label="REVIEW_EVENT\nchapter_id FK; revision_id FK\nactor_id FK; action"];
          chapter->revision; revision->comment;
          chapter->event; revision->event [style=dashed];
          user->revision; user->comment; user->event;
          {rank=same; comment; event;}
        }''',
        6: r'''digraph G {''' + base + r'''
          rankdir=TB; splines=polyline;
          project [label="PROJECT\nid UUID PK"];
          row [label="CONSISTENCY_ROW\nproject_id + ordinal UK\nquestion; objective; hypothesis"];
          variable [label="RESEARCH_VARIABLE\nproject_id FK; role; name"];
          dimension [label="VARIABLE_DIMENSION\nvariable_id + ordinal UK"];
          indicator [label="INDICATOR\ndimension_id + ordinal UK\nmeasurement_scale"];
          instrument [label="INSTRUMENT\nproject_id + lineage_key + version_no UK"];
          item [label="INSTRUMENT_ITEM\ninstrument_id + ordinal UK\ndimension_id FK"];
          project->row; project->variable; variable->dimension->indicator;
          project->instrument->item; dimension->item [style=dashed];
          {rank=same; row; variable; instrument;}
          {rank=same; indicator; item;}
        }''',
        7: r'''digraph G {''' + base + r'''
          rankdir=TB; splines=polyline;
          project [label="PROJECT\nid UUID PK"];
          source [label="SOURCE_DOCUMENT\nproject_id FK; kind\nverification_status"];
          excerpt [label="SOURCE_EXCERPT\nsource_id + chunk_no UK\npage_no; text"];
          revision [label="CHAPTER_REVISION\nid PK; chapter_id FK"];
          citation [label="EVIDENCE_CITATION\nrevision_id FK; excerpt_id FK\nclaim_start_cp; claim_end_cp"];
          project->source->excerpt->citation;
          revision->citation;
        }''',
    }
    if which not in diagrams:
        raise ValueError(f"No hay render de diagrama {which}")
    path = TEMP / f"diagrama_{which}.png"
    DIAGRAM_SOURCE_OUT.mkdir(exist_ok=True)
    (DIAGRAM_SOURCE_OUT / f"figura_{which}_word.dot").write_text(diagrams[which] + "\n", encoding="utf-8")
    subprocess.run(["dot", "-Tpng", "-Gdpi=170", "-o", str(path)], input=diagrams[which], text=True, check=True)
    return path


def er_diagram_png(
    number: int, mermaid_lines: list[str],
    selected: set[str] | None = None, suffix: str = "",
) -> Path:
    """Render the ER source in the report, including every persisted column."""
    if not mermaid_lines or mermaid_lines[0].strip() != "erDiagram":
        raise ValueError(f"La figura {number} no es un diagrama ER")
    entities: dict[str, list[tuple[str, str, str]]] = {}
    relations: list[tuple[str, str, str, str, str]] = []
    current = None
    relation_pattern = re.compile(
        r"^([A-Z_]+)\s+([|o}{]+)--([|o}{]+)\s+([A-Z_]+)\s*:\s*(\S+)"
    )
    for raw in mermaid_lines[1:]:
        line = raw.strip()
        if not line or line.startswith("%%"):
            continue
        if line == "}":
            current = None
        elif line.endswith("{"):
            current = line[:-1].strip()
            entities[current] = []
        elif current:
            parts = line.split()
            if len(parts) < 2:
                raise ValueError(f"Atributo ER inválido: {line}")
            entities[current].append((parts[0], parts[1], " ".join(parts[2:])))
        else:
            match = relation_pattern.match(line)
            if not match:
                raise ValueError(f"Relación ER inválida: {line}")
            relations.append(match.groups())
    if any(a not in entities or b not in entities for a, _, _, b, _ in relations):
        raise ValueError(f"La figura {number} tiene una relación sin entidad")
    if selected is not None:
        entities = {name: fields for name, fields in entities.items() if name in selected}
        relations = [r for r in relations if r[0] in selected and r[3] in selected]

    def cardinality(mark: str) -> str:
        return {"||": "1", "o|": "0..1", "|{": "1..N", "o{": "0..N"}[mark]

    palette = {"PROJECT": "#17324D", "USER": "#17324D"}
    lines = [
        "digraph ER {",
        f'graph [bgcolor="white", rankdir={"LR" if number == 8 else "TB"}, splines=polyline, '
        'nodesep=0.36, ranksep=0.68, pad=0.12, outputorder=edgesfirst];',
        'node [shape=plain, fontname="Arial"];',
        'edge [color="#64748B", fontname="Arial", fontsize=9, arrowsize=0.55];',
    ]
    for name, fields in entities.items():
        color = palette.get(name, "#1E5A78")
        rows = [
            '<TR><TD COLSPAN="3" BGCOLOR="' + color + '">'
            '<FONT COLOR="white"><B>' + escape(name) + '</B></FONT></TD></TR>'
        ]
        for typ, field, keys in fields:
            rows.append(
                '<TR><TD ALIGN="LEFT"><FONT COLOR="#475569">' + escape(typ) + '</FONT></TD>'
                '<TD ALIGN="LEFT">' + escape(field) + '</TD>'
                '<TD ALIGN="LEFT">' + ("<B>" + escape(keys) + "</B>" if keys else "&#160;") + '</TD></TR>'
            )
        label = (
            '<<TABLE BORDER="1" CELLBORDER="0" CELLSPACING="0" CELLPADDING="4" '
            'COLOR="#CBD5E1">' + "".join(rows) + "</TABLE>>"
        )
        lines.append(f'"{name}" [label={label}];')
    for source, left, right, target, name in relations:
        cardinalities = f"{cardinality(left)} : {cardinality(right)}"
        lines.append(
            f'"{source}" -> "{target}" '
            f'[label="{escape(name)}\\n{cardinalities}", minlen=1];'
        )
    lines.append("}")
    path = TEMP / f"diagrama_{number}{suffix}_er.png"
    DIAGRAM_SOURCE_OUT.mkdir(exist_ok=True)
    (DIAGRAM_SOURCE_OUT / f"figura_{number}{suffix}_word.dot").write_text("\n".join(lines) + "\n", encoding="utf-8")
    subprocess.run(
        ["dot", "-Tpng", "-Gdpi=190", "-o", str(path)],
        input="\n".join(lines), text=True, check=True,
    )
    return path


def add_diagram_picture(
    doc: Document, path: Path, caption_text: str, landscape: bool = False,
    height_limit_cm: float | None = None,
) -> None:
    with Image.open(path) as im:
        width_px, height_px = im.size
    max_width, max_height = (12.8, 16.0) if landscape else (8.65, 23.5)
    if height_limit_cm is not None:
        max_height = height_limit_cm
    display_width = min(max_width, max_height * width_px / height_px)
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_together = True
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(path), width=Cm(display_width))
    caption = doc.add_paragraph(caption_text)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(8)


def add_diagram(doc: Document, number: int, mermaid_lines: list[str]) -> None:
    labels = {
        1: "Figura 1. Flujo operativo propuesto para tesista y asesor.",
        2: "Figura 2. Secuencia de preparación y guardado de un paso.",
        3: "Figura 3. Estados permitidos en la revisión de un capítulo.",
        4: "Figura 4. Identidad, proyecto y avance metodológico.",
        5: "Figura 5. Versiones, comentarios y auditoría de revisión.",
        6: "Figura 6. Matriz de consistencia, variables e indicadores.",
        7: "Figura 7. Instrumentos y trazabilidad de indicadores.",
        8: "Figura 8. Fuentes, extractos y evidencia citada.",
    }
    if number == 4:
        selected_a = {"USER", "PROJECT", "PROJECT_PARTICIPANT", "PROJECT_INVITATION"}
        selected_b = {"PROJECT", "RESEARCH_PROFILE", "STEP_PROGRESS", "CHAPTER"}
        add_diagram_picture(
            doc, er_diagram_png(number, mermaid_lines, selected_a, "a"),
            "Figura 4a. Identidad y participación en el proyecto.", True,
        )
        add_diagram_picture(
            doc, er_diagram_png(number, mermaid_lines, selected_b, "b"),
            "Figura 4b. Perfil, pasos y capítulos del proyecto.", True,
        )
    elif number == 5:
        add_diagram_picture(doc, er_diagram_png(number, mermaid_lines), labels[number], True)
    elif number == 8:
        add_diagram_picture(doc, er_diagram_png(number, mermaid_lines), labels[number], True)
    else:
        path = diagram_png(number) if number <= 3 else er_diagram_png(number, mermaid_lines)
        add_diagram_picture(
            doc, path, labels[number],
            landscape=number >= 6,
            height_limit_cm=13.5 if number == 2 else 14.7 if number in (6, 7) else None,
        )


def add_table(doc: Document, rows: list[str], compact: bool = False) -> None:
    cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    cells = [cells[0]] + cells[2:]  # discard Markdown alignment row
    count = len(cells[0])
    table = doc.add_table(rows=1, cols=count)
    table.autofit = False
    table.alignment = 1
    if compact and count == 4 and cells[0][0] == "Entidad":
        widths = [Cm(2.4), Cm(2.7), Cm(4.6), Cm(2.7)]
        size = 7.6
    elif compact and count == 4:
        widths = [Cm(2.2), Cm(3.7), Cm(2.6), Cm(3.9)]
        size = 7.2
    elif count >= 10:
        widths = [Cm(5.3)] + [Cm(1.14)] * (count - 1)
        size = 6.5
    elif count >= 5:
        widths = [Cm(2.8)] + [Cm(3.1)] * (count - 1)
        size = 8.0
    elif count == 4 and cells[0][0] == "Entidad":
        widths = [Cm(3.0), Cm(3.6), Cm(5.4), Cm(3.4)]
        size = 8.5
    elif count == 4:
        widths = [Cm(2.5), Cm(4.6), Cm(2.9), Cm(5.4)]
        size = 8.5
    else:
        if count == 2 and "Factores Internos" in cells[0][0]:
            widths = [Cm(7.7), Cm(7.7)]
        else:
            widths = [Cm(3.7), Cm(11.7)] if count == 2 else [Cm(5.0)] * count
        size = 8.8
    for j, width in enumerate(widths):
        table.columns[j].width = width
    for r, values in enumerate(cells):
        row = table.rows[0] if r == 0 else table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        if r == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for j, raw in enumerate(values):
            cell = row.cells[j]
            if j < len(widths):
                cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            set_cell_margin(cell, 45 if count >= 5 or compact else 80)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.line_spacing = 1.08
            if count >= 10 and j > 0:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            active = count >= 10 and "████" in raw
            add_formatted(para, "" if active else raw)
            for run in para.runs:
                run.font.size = Pt(size)
                if r == 0:
                    run.bold = True
            if r == 0:
                set_cell_shading(cell, "E9EDF1")
            elif active:
                set_cell_shading(cell, "D9E6F2")
            elif r % 2 == 0:
                set_cell_shading(cell, "F8F9FA")
            set_cell_border(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def set_section(section, landscape: bool, columns: int = 1, frontmatter: bool = False) -> None:
    section.page_width, section.page_height = (Cm(29.7), Cm(21.0)) if landscape else (Cm(21.0), Cm(29.7))
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.top_margin = Cm(2.0 if frontmatter else 1.7)
    section.bottom_margin = Cm(2.0 if frontmatter else 1.7)
    section.left_margin = Cm(2.5 if frontmatter else 1.5)
    section.right_margin = Cm(2.5 if frontmatter else 1.5)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)
    cols = section._sectPr.xpath("./w:cols")
    cols = cols[0] if cols else OxmlElement("w:cols")
    cols.set(qn("w:num"), str(columns))
    cols.set(qn("w:space"), "340")
    if cols.getparent() is None:
        section._sectPr.append(cols)


def build_docx(markdown: str, output_path: Path | str = DOCX_OUT) -> Path:
    doc = Document()
    set_section(doc.sections[0], False, 1)
    styles = doc.styles
    set_font(styles["Normal"], 10)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Normal"].paragraph_format.space_after = Pt(3)
    for style_name, size in [("Title", 16), ("Heading 1", 11), ("Heading 2", 10.5), ("Heading 3", 10), ("Heading 4", 10)]:
        set_font(styles[style_name], size, True)
        styles[style_name].paragraph_format.space_before = Pt(9)
        styles[style_name].paragraph_format.space_after = Pt(4)
        styles[style_name].paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(14)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    # A simple page number field is easier to maintain than a static count.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    lines = markdown.splitlines()
    i = 0
    diagram = 0
    in_index = False
    in_landscape = False
    in_er_landscape = False
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith("```mermaid"):
            diagram += 1
            if diagram == 4:
                set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), False, 1)
                set_section(doc.add_section(WD_SECTION_START.NEW_PAGE), True, 2)
                in_er_landscape = True
            i += 1
            source = []
            while i < len(lines) and lines[i] != "```":
                source.append(lines[i])
                i += 1
            if i == len(lines):
                raise ValueError(f"Diagrama Mermaid {diagram} sin cierre")
            add_diagram(doc, diagram, source)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            landscape = len(rows[0].strip("|").split("|")) >= 10
            if in_er_landscape:
                add_table(doc, rows, compact=True)
                continue
            if landscape and not in_landscape:
                section = doc.add_section(WD_SECTION_START.CONTINUOUS)
                set_section(section, False, 1)
            elif not landscape:
                set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), False, 1)
            add_table(doc, rows)
            if landscape:
                section = doc.add_section(WD_SECTION_START.CONTINUOUS)
                set_section(section, False, 2)
                in_landscape = False
            else:
                set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), False, 2)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_formatted(p, line[2:])
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:]
            if title == "ÍNDICE DEL INFORME":
                p = doc.add_paragraph(style="Heading 1")
                add_formatted(p, title)
                j = i + 1
                entries = []
                while j < len(lines) and not lines[j].startswith("## "):
                    if lines[j].lstrip().startswith("*"):
                        entries.append(lines[j])
                    j += 1
                split = next(k for k, entry in enumerate(entries) if "CAPÍTULO IV:" in entry)
                table = doc.add_table(rows=1, cols=2)
                table.autofit = False
                for col, group in enumerate((entries[:split], entries[split:])):
                    cell = table.cell(0, col)
                    cell.width = Cm(8.7)
                    set_cell_margin(cell, 35)
                    cell.text = ""
                    for k, entry in enumerate(group):
                        para = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                        stripped = entry.lstrip()
                        depth = (len(entry) - len(stripped)) // 2
                        para.paragraph_format.left_indent = Cm(0.22 + depth * 0.42)
                        para.paragraph_format.first_line_indent = Cm(-0.20)
                        para.paragraph_format.space_after = Pt(1)
                        para.paragraph_format.line_spacing = 1.0
                        add_formatted(para, "• " + stripped[2:])
                        for run in para.runs:
                            run.font.size = Pt(8.5)
                i = j
                continue
            elif title.startswith("CAPÍTULO I:"):
                set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), False, 2)
            p = doc.add_paragraph(style="Heading 1")
            add_formatted(p, title)
            i += 1
            continue
        if line.startswith("### "):
            if line.startswith("### 4.3 Cronograma"):
                section = doc.add_section(WD_SECTION_START.CONTINUOUS)
                set_section(section, False, 1)
                in_landscape = True
            p = doc.add_paragraph(style="Heading 2")
            add_formatted(p, line[4:])
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_formatted(p, line[5:])
            i += 1
            continue
        if re.match(r"^\[\d+\] ", line):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            add_formatted(p, line)
            i += 1
            continue
        stripped = line.lstrip()
        if re.match(r"^(?:\*|\d+\.)\s", stripped):
            indent = (len(line) - len(stripped)) // 2
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Cm(0.55 + indent * 0.45)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_after = Pt(2)
            item = re.sub(r"^(?:\*|\d+\.)\s+", "", stripped)
            marker = "• " if stripped.startswith("*") else re.match(r"^\d+\.", stripped).group(0) + " "
            add_formatted(p, marker + item)
            if re.search(r"\[[1-7]\]", item):
                p.paragraph_format.keep_with_next = True
            i += 1
            continue
        p = doc.add_paragraph(style="Normal")
        if in_index:
            p.paragraph_format.left_indent = Cm(0.4)
        add_formatted(p, line)
        i += 1

    # A final continuous section makes Word balance the last two columns.
    set_section(doc.add_section(WD_SECTION_START.CONTINUOUS), True, 1)
    doc.core_properties.title = "Informe del proyecto Curso Integrador I"
    doc.core_properties.subject = "TesisMCP"
    doc.save(str(output_path))
    return Path(output_path)


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Compilador de informes Markdown a DOCX con formato IEEE.")
    parser.add_argument("--input", "-i", type=Path, default=MD_OUT, help="Ruta al archivo Markdown fuente.")
    parser.add_argument("--output", "-o", type=Path, default=DOCX_OUT, help="Ruta al archivo DOCX de salida.")
    args = parser.parse_args()

    content = args.input.read_text(encoding="utf-8")
    out_file = build_docx(content, args.output)
    print(f"Fuente Markdown: {args.input}")
    print(f"Documento generado: {out_file}")
